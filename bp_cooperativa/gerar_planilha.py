"""Gera a planilha com o BP de exemplo, os indicadores e os exercícios resolvidos.

Tudo que é conta vira fórmula do Excel, ligada às abas BP, DSP e Dados.
Convenção de cores: azul = dado digitado, preto = fórmula, verde = link de outra aba.
"""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import modelo as M
from exercicios import CLASSIFICACAO, EXERCICIOS

S, D, I, NOTAS = M.S, M.D, M.I, M.NOTAS
SAIDA = "Exercicios_BP_Cooperativa_Resolvidos.xlsx"

FONTE = "Arial"
AZUL, VERDE, PRETO, BRANCO = "0000FF", "008000", "000000", "FFFFFF"
COR_TITULO = "1F4E5A"
COR_SECAO = "D9E7EA"
COR_SUB = "F2F2F2"
COR_RESPOSTA = "FFF2CC"

F_MIL = '#,##0;(#,##0);"-"'
F_MIL1 = '#,##0.0;(#,##0.0);"-"'
F_MIL2 = '#,##0.00;(#,##0.00);"-"'
F_PCT1 = '0.0%;-0.0%;"-"'
F_PCT2 = '0.00%;-0.00%;"-"'
F_PP = '0.00" p.p.";-0.00" p.p.";"-"'
F_X = '0.00"x"'
F_RS = '"R$ "#,##0.00'

FMT_UNIDADE = {"%": F_PCT2, "R$ mil": F_MIL, "p.p.": F_PP, "x": F_X, "R$": F_RS}

fino = Side(style="thin", color="BFBFBF")
grosso = Side(style="medium", color="1F4E5A")

REG = {}  # (aba, chave) -> linha
COL_ANO = {
    "BP": {2024: "B", 2025: "C"},
    "DSP": {2024: "B", 2025: "C"},
    "Gerencial": {2024: "B", 2025: "C"},
    "Dados": {2023: "B", 2024: "C", 2025: "D"},
    "Indicadores": {2024: "C", 2025: "D"},
}


def ref(aba, chave, ano=None, col=None):
    linha = REG[(aba, chave)]
    c = col or COL_ANO[aba][ano]
    return f"{aba}!${c}${linha}"


def bp(k, a):
    return ref("BP", k, a)


def ds(k, a):
    return ref("DSP", k, a)


def ge(k, a):
    return ref("Gerencial", k, a)


def da(k, a):
    return ref("Dados", k, a)


def par(k):
    return ref("Dados", k, col="B")


def media(k, a):
    return f"AVERAGE({da(k, a - 1)},{da(k, a)})"


def rw(k, a):
    linha = REG[("RWA", k)]
    return f"RWA!${'D' if a == 2024 else 'F'}${linha}"


# ------------------------------------------------------------------ tradução PT-BR
FUNCOES = {"AVERAGE": "MÉDIA", "SUM": "SOMA", "MIN": "MÍNIMO", "MAX": "MÁXIMO", "IFERROR": "SEERRO",
           "ROUND": "ARRED", "IF": "SE", "ABS": "ABS", "SUMPRODUCT": "SOMARPRODUTO"}


def pt(formula):
    if not isinstance(formula, str) or not formula.startswith("="):
        return ""
    f = formula.replace("$", "")
    for en in sorted(FUNCOES, key=len, reverse=True):
        f = re.sub(rf"\b{en}\(", FUNCOES[en] + "(", f)
    out, aspas = [], False
    for ch in f:
        if ch == '"':
            aspas = not aspas
        out.append(";" if ch == "," and not aspas else ch)
    return re.sub(r"(?<![A-Za-z0-9$])(\d+)\.(\d+)", r"\1,\2", "".join(out))


# ------------------------------------------------------------------ estilos
def estilo(c, cor=PRETO, negrito=False, fmt=None, fundo=None, alinhar=None, italico=False, tam=10, quebra=False):
    c.font = Font(name=FONTE, size=tam, bold=negrito, color=cor, italic=italico)
    if fmt:
        c.number_format = fmt
    if fundo:
        c.fill = PatternFill("solid", fgColor=fundo)
    if alinhar or quebra:
        c.alignment = Alignment(horizontal=alinhar, vertical="top", wrap_text=quebra)


def cor_de(valor):
    if isinstance(valor, str) and valor.startswith("="):
        corpo = valor[1:]
        if re.fullmatch(r"[A-Za-z]+!\$?[A-Z]+\$?\d+", corpo):
            return VERDE
        return PRETO
    return AZUL


def titulo(ws, texto, sub=None, largura=8):
    ws["A1"] = texto
    estilo(ws["A1"], BRANCO, True, fundo=COR_TITULO, tam=13)
    for c in range(2, largura + 1):
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=COR_TITULO)
    if sub:
        ws["A2"] = sub
        estilo(ws["A2"], "595959", italico=True, tam=9)


def cabecalho(ws, linha, nomes):
    for j, n in enumerate(nomes, start=1):
        c = ws.cell(row=linha, column=j, value=n)
        estilo(c, BRANCO, True, fundo=COR_TITULO, alinhar="center", quebra=True)


def larguras(ws, ls):
    for j, w in enumerate(ls, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


wb = Workbook()
wb.remove(wb.active)

# ================================================================== LEIA-ME
ws = wb.create_sheet("LEIA-ME")
titulo(ws, "Análise de Balanço de Cooperativa de Crédito: planilha de exercícios resolvidos", None, 3)
textos = [
    ("Para que serve", "Esta planilha acompanha o material HTML 'Como analisar o BP de uma cooperativa de crédito'. Ela traz o BP e a DSP de exemplo, as notas, todos os indicadores e a resolução de cada exercício proposto, com fórmula e explicação."),
    ("Cooperativa de exemplo", "Coop Vale do Ipê: cooperativa singular fictícia, filiada a uma central de um sistema cooperativo, segmento S4. Os números foram construídos para fins didáticos e fecham entre si (o balanço bate, a DSP explica o PL, os saldos médios explicam as receitas)."),
    ("Unidade", "Valores em R$ mil, salvo indicação. Ex.: 1.870.000 = R$ 1,87 bilhão."),
    ("Cores", "Azul = dado digitado (pode mudar para simular). Preto = fórmula. Verde = link direto de outra aba. Amarelo = resposta do exercício."),
    ("Fórmulas", "A coluna 'Fórmula (Excel PT-BR)' das abas de exercício mostra a fórmula como ela aparece no Excel em português (MÉDIA, SOMA, MÍNIMO, ponto e vírgula). O Excel converte automaticamente."),
    ("Ordem sugerida", "1) BP, DSP e Gerencial para conhecer a cooperativa. 2) Dados (notas). 3) Indicadores. 4) Ex01 a Ex15, na ordem do HTML. 5) Gabarito para conferir tudo de uma vez."),
    ("Abas", "BP | DSP | Gerencial (DRE em cascata) | Dados (notas e premissas) | DMPL | RWA | Indicadores | Ex01 a Ex15 | Gabarito"),
    ("Ressalvas", "Os fatores de ponderação (FPR), o fator do risco operacional e as premissas de estresse do LCR são simplificações didáticas. Na cooperativa real, confira no DLO (documento 2061), no demonstrativo de RWA e na política de liquidez. A régua de 10,5% (8% + 2,5% de conservação) vale para o S4."),
    ("Simule", "Troque um número azul (por exemplo, a despesa de PDD de 2025 na aba DSP) e veja ROE, IEAR, PL e Basileia mudarem juntos. É o melhor jeito de entender como as caixas se ligam."),
]
for n, (k, v) in enumerate(textos, start=3):
    ws.cell(row=n, column=1, value=k)
    estilo(ws.cell(row=n, column=1), PRETO, True)
    ws.cell(row=n, column=2, value=v)
    estilo(ws.cell(row=n, column=2), quebra=True)
    ws.row_dimensions[n].height = 48
larguras(ws, [24, 120, 4])

# ================================================================== BP
DESC_BP = {
    "caixa": "Dinheiro em espécie nas agências e caixas eletrônicos.",
    "centralizacao": "Sobras de recursos aplicadas na Cooperativa Central. Principal reserva de liquidez; rende perto de 100% do CDI.",
    "aplic_interfin": "Depósitos interfinanceiros aplicados (DI intrassistema, Banco Cooperativo) e compromissadas.",
    "tvm": "Títulos públicos federais e cotas de fundos de liquidez.",
    "cred_emprestimos": "Crédito pessoal, consignado, cheque especial, capital de giro, desconto de títulos, cartão parcelado.",
    "cred_financiamentos": "Financiamento de veículos, imóveis, máquinas e investimentos de PJ (inclui repasses BNDES).",
    "cred_rural": "Custeio, investimento e comercialização agrícola e pecuária; agroindústria.",
    "carteira": "Soma das operações de crédito antes da provisão.",
    "provisao": "Perda esperada da carteira (Res. CMN 4.966). Conta redutora: negativa.",
    "carteira_liq": "Carteira bruta menos provisão: o que se espera receber.",
    "outros_ativos_fin": "Rendas a receber, valores de cartões a receber, convênios, devedores diversos financeiros.",
    "ativos_fiscais": "Impostos a compensar e créditos tributários.",
    "bndu": "Bens recebidos em dação em pagamento, à venda.",
    "outros_ativos": "Despesas antecipadas, adiantamentos, materiais.",
    "investimentos": "Quotas da Central, da holding do sistema e de outras entidades.",
    "imobilizado": "Agências, terrenos, móveis, equipamentos e TI, líquidos de depreciação.",
    "intangivel": "Softwares e licenças. Deduzido do capital no PR.",
    "ativo": "Tudo o que a cooperativa tem ou tem a receber.",
    "dep_vista": "Saldo de conta corrente. Não paga juros; sai a qualquer momento.",
    "poupanca": "Poupança (inclui a rural). Remuneração da regra da poupança; parte é direcionada ao crédito rural.",
    "rdc": "Recibo de Depósito Cooperativo: depósito a prazo pago em % do CDI.",
    "depositos": "Depósitos totais dos associados.",
    "lca": "Letra de Crédito do Agronegócio emitida pela cooperativa.",
    "repasses_interfin": "Recursos tomados da Central e do Banco do sistema (inclui recursos para crédito rural e recursos corporativos).",
    "emprestimos_repasses": "BNDES, FINAME, BRDE, Funcafé e outros fundos repassados.",
    "lfs": "Letra Financeira Subordinada: dívida longa que conta como capital (Nível II).",
    "outros_passivos_fin": "Cobrança e arrecadação de terceiros, valores a repassar de cartões.",
    "provisoes": "Contingências trabalhistas, cíveis e fiscais e provisão para garantias financeiras prestadas.",
    "obrig_sociais": "FATES, juros ao capital e sobras a pagar, quotas a devolver de associados desligados.",
    "obrig_fiscais_outras": "Tributos a recolher, salários e encargos, fornecedores.",
    "passivo": "Obrigações com associados e terceiros.",
    "capital_social": "Quotas-partes integralizadas pelos associados. Podem ser devolvidas quando o associado sai.",
    "reserva_legal": "Parte das sobras retida por lei e estatuto. Capital mais estável.",
    "sobras_acum": "Sobras do ano ainda sem destino: a AGO decide (capital, conta corrente, reservas).",
    "aap": "Marcação a mercado de títulos classificados em VJORA.",
    "pl": "Recursos próprios dos associados.",
    "coobrigacoes": "Fianças e coobrigações em repasses. Não está no ativo, mas consome capital.",
}

LINHAS_BP = [
    ("sec", None, "ATIVO"),
    ("in", "caixa", "Disponibilidades (caixa)"),
    ("in", "centralizacao", "Centralização financeira (relações interfinanceiras)"),
    ("in", "aplic_interfin", "Aplicações interfinanceiras de liquidez"),
    ("in", "tvm", "Títulos e valores mobiliários"),
    ("in", "cred_emprestimos", "Operações de crédito: empréstimos e títulos descontados"),
    ("in", "cred_financiamentos", "Operações de crédito: financiamentos"),
    ("in", "cred_rural", "Operações de crédito: financiamentos rurais e agroindustriais"),
    ("sub", "carteira", "(=) Carteira de crédito bruta", ["cred_emprestimos", "cred_financiamentos", "cred_rural"]),
    ("in", "provisao", "(-) Provisão para perdas esperadas (risco de crédito)"),
    ("sub", "carteira_liq", "(=) Carteira de crédito líquida", ["carteira", "provisao"]),
    ("in", "outros_ativos_fin", "Outros ativos financeiros"),
    ("in", "ativos_fiscais", "Ativos fiscais (correntes e diferidos)"),
    ("in", "bndu", "Bens não de uso próprio (BNDU)"),
    ("in", "outros_ativos", "Outros ativos"),
    ("in", "investimentos", "Investimentos (participações no sistema)"),
    ("in", "imobilizado", "Imobilizado de uso"),
    ("in", "intangivel", "Intangível"),
    ("tot", "ativo", "TOTAL DO ATIVO", ["caixa", "centralizacao", "aplic_interfin", "tvm", "carteira_liq", "outros_ativos_fin",
                                        "ativos_fiscais", "bndu", "outros_ativos", "investimentos", "imobilizado", "intangivel"]),
    ("sec", None, "PASSIVO"),
    ("in", "dep_vista", "Depósitos à vista"),
    ("in", "poupanca", "Depósitos de poupança"),
    ("in", "rdc", "Depósitos a prazo (RDC)"),
    ("sub", "depositos", "(=) Total de depósitos", ["dep_vista", "poupanca", "rdc"]),
    ("in", "lca", "Recursos de aceites e emissão de títulos (LCA)"),
    ("in", "repasses_interfin", "Relações interfinanceiras (repasses da Central e do Banco)"),
    ("in", "emprestimos_repasses", "Obrigações por empréstimos e repasses (BNDES e outros)"),
    ("in", "lfs", "Instrumentos de dívida elegíveis a capital (LFS)"),
    ("in", "outros_passivos_fin", "Outros passivos financeiros"),
    ("in", "provisoes", "Provisões (contingências e garantias prestadas)"),
    ("in", "obrig_sociais", "Obrigações sociais e estatutárias (FATES, sobras e quotas a pagar)"),
    ("in", "obrig_fiscais_outras", "Obrigações fiscais e outras obrigações"),
    ("tot", "passivo", "TOTAL DO PASSIVO", ["depositos", "lca", "repasses_interfin", "emprestimos_repasses", "lfs",
                                            "outros_passivos_fin", "provisoes", "obrig_sociais", "obrig_fiscais_outras"]),
    ("sec", None, "PATRIMÔNIO LÍQUIDO"),
    ("in", "capital_social", "Capital social"),
    ("in", "reserva_legal", "Reserva legal"),
    ("in", "sobras_acum", "Sobras acumuladas (à disposição da AGO)"),
    ("in", "aap", "Ajustes de avaliação patrimonial"),
    ("tot", "pl", "TOTAL DO PATRIMÔNIO LÍQUIDO", ["capital_social", "reserva_legal", "sobras_acum", "aap"]),
    ("tot", "passivo_pl", "TOTAL DO PASSIVO + PL", ["passivo", "pl"]),
    ("chk", "check", "Conferência: ativo - (passivo + PL) deve ser zero", None),
    ("sec", None, "FORA DO BALANÇO (contas de compensação)"),
    ("in", "coobrigacoes", "Coobrigações e garantias prestadas"),
]
ATIVO_KEYS = {"caixa", "centralizacao", "aplic_interfin", "tvm", "cred_emprestimos", "cred_financiamentos", "cred_rural",
              "carteira", "provisao", "carteira_liq", "outros_ativos_fin", "ativos_fiscais", "bndu", "outros_ativos",
              "investimentos", "imobilizado", "intangivel", "ativo"}


def valor_bp(k, a):
    if k == "coobrigacoes":
        return NOTAS["coobrigacoes"][a]
    return S[k][a]


def montar_demonstrativo(ws, linhas, base_fn, inicio, aba, valor_fn, descr=None, av_base_nome=""):
    r = inicio
    for ln in linhas:
        tipo, k, rot = ln[0], ln[1], ln[2]
        if tipo == "sec":
            ws.cell(row=r, column=1, value=rot)
            for c in range(1, 9):
                estilo(ws.cell(row=r, column=c), COR_TITULO, True, fundo=COR_SECAO)
            ws.cell(row=r, column=1).value = rot
            r += 1
            continue
        REG[(aba, k)] = r
        r += 1
    # segunda passada: escrever valores (precisa dos endereços)
    r = inicio
    for ln in linhas:
        tipo, k, rot = ln[0], ln[1], ln[2]
        if tipo == "sec":
            r += 1
            continue
        ws.cell(row=r, column=1, value=rot)
        negrito = tipo in ("tot", "sub", "fx_tot")
        fundo = COR_SUB if tipo in ("sub", "tot") else None
        for a, col in COL_ANO[aba].items():
            if tipo == "in":
                v = valor_fn(k, a)
            elif tipo in ("sub", "tot"):
                v = "=" + "+".join(f"{col}{REG[(aba, x)]}" for x in ln[3])
            elif tipo == "chk":
                v = f"={col}{REG[(aba, 'ativo')]}-{col}{REG[(aba, 'passivo_pl')]}"
            elif tipo == "fx":
                v = ln[3](col)
            c = ws.cell(row=r, column="ABCDEFGH".index(col) + 1, value=v)
            estilo(c, cor_de(v), negrito, F_MIL, fundo)
        estilo(ws.cell(row=r, column=1), PRETO, negrito, fundo=fundo)
        if tipo != "chk":
            base24, base25 = base_fn(k)
            if base24:
                for col_av, col_v, base in (("D", "B", base24), ("E", "C", base25)):
                    c = ws.cell(row=r, column="ABCDEFGH".index(col_av) + 1, value=f"=IFERROR({col_v}{r}/{col_v}${base},0)")
                    estilo(c, PRETO, negrito, F_PCT1, fundo)
            c = ws.cell(row=r, column=6, value=f'=IFERROR(C{r}/B{r}-1,"n.a.")')
            estilo(c, PRETO, negrito, F_PCT1, fundo, "right")
            c = ws.cell(row=r, column=7, value=f"=C{r}-B{r}")
            estilo(c, PRETO, negrito, F_MIL, fundo)
            if descr and k in descr:
                c = ws.cell(row=r, column=8, value=descr[k])
                estilo(c, "404040", tam=9, fundo=fundo)
        r += 1
    return r


ws = wb.create_sheet("BP")
titulo(ws, "Balanço Patrimonial: Coop Vale do Ipê (fictícia)", "Em R$ mil | posição em 31/12 | AV sobre o total do ativo (ou do passivo + PL) | AH = 2025 ÷ 2024 - 1")
cabecalho(ws, 4, ["Linha", "31/12/2024", "31/12/2025", "AV 2024", "AV 2025", "AH 25/24", "Variação (R$ mil)", "O que entra"])
# registrar linhas primeiro para as bases da AV
r0 = 5
_r = r0
for ln in LINHAS_BP:
    if ln[0] != "sec":
        REG[("BP", ln[1])] = _r
    _r += 1


def base_bp(k):
    if k in ("check", "coobrigacoes"):
        return (None, None)
    b = REG[("BP", "ativo")] if k in ATIVO_KEYS else REG[("BP", "passivo_pl")]
    return (b, b)


fim_bp = montar_demonstrativo(ws, LINHAS_BP, base_bp, r0, "BP", valor_bp, DESC_BP)
larguras(ws, [58, 14, 14, 9, 9, 10, 14, 90])
ws.freeze_panes = "B5"

# ================================================================== DSP
DESC_DSP = {
    "rend_credito": "Juros e encargos da carteira (comercial e rural). Ativo problemático (estágio 3) não gera receita até ser recebido.",
    "recuperacao": "Valores recebidos de créditos que já tinham sido baixados como perda.",
    "rend_centralizacao": "Rendimento da centralização financeira na Central.",
    "rend_aplic": "Rendimento do DI aplicado e de outras aplicações interfinanceiras.",
    "rend_tvm": "Rendimento dos títulos e fundos.",
    "desp_poupanca": "Remuneração paga na poupança.",
    "desp_rdc": "Juros pagos (ou apropriados) no RDC.",
    "desp_lca": "Juros pagos na LCA.",
    "desp_lfs": "Juros da LFS.",
    "fgcoop_neg": "Contribuição ordinária ao Fundo Garantidor do Cooperativismo de Crédito.",
    "desp_repasses_interfin": "Custo dos recursos tomados da Central e do Banco.",
    "desp_emprestimos_repasses": "Custo dos repasses BNDES e outros fundos.",
    "pdd": "Despesa de provisão para perdas esperadas do período.",
    "servicos": "Cartões, seguros, consórcios, cobrança, comissões.",
    "tarifas": "Tarifas de conta e pacotes de serviços.",
    "pessoal": "Salários, encargos, benefícios, PLR e treinamento.",
    "outras_adm": "TI, aluguéis, depreciação, marketing, transporte de valores, rateios da Central e da Confederação.",
    "tributarias": "PIS, Cofins e ISS (principalmente sobre atos não cooperativos e serviços).",
    "outras_rec_op": "Juros ao capital recebidos de participações, recuperação de encargos e outras.",
    "outras_desp_op": "Despesas com cartões, perdas operacionais, descontos concedidos e outras.",
    "contingencias": "Constituição (-) ou reversão (+) de provisões para contingências e garantias.",
    "nao_operacional": "Ganho ou perda na venda de bens (imóveis, BNDU).",
    "ir_cs": "Tributos sobre o resultado de atos não cooperativos.",
    "resultado": "Base do ROE: resultado gerado antes de remunerar as quotas.",
    "juros_capital": "Remuneração das quotas-partes (limitada à Selic pela LC 130/2009).",
    "dest_fates_ato": "Resultado de atos não cooperativos vai inteiro para o FATES (Lei 5.764, art. 87).",
    "dest_reserva": "Percentual do estatuto sobre a base (lei: mínimo de 10%).",
    "dest_fates": "Percentual do estatuto sobre a base (lei: mínimo de 5%).",
    "sobras_ago": "A AGO decide: capitalizar, distribuir em conta ou reforçar reservas.",
}
LINHAS_DSP = [
    ("sec", None, "RECEITAS (INGRESSOS) DA INTERMEDIAÇÃO FINANCEIRA"),
    ("in", "rend_credito", "Operações de crédito"),
    ("in", "recuperacao", "Recuperação de créditos baixados como prejuízo"),
    ("in", "rend_centralizacao", "Ingressos de depósitos intercooperativos (centralização financeira)"),
    ("in", "rend_aplic", "Resultado de aplicações interfinanceiras de liquidez"),
    ("in", "rend_tvm", "Resultado de operações com títulos e valores mobiliários"),
    ("sub", "receitas_intermed", "(=) Total das receitas da intermediação", ["rend_credito", "recuperacao", "rend_centralizacao", "rend_aplic", "rend_tvm"]),
    ("sec", None, "DESPESAS (DISPÊNDIOS) DA INTERMEDIAÇÃO FINANCEIRA"),
    ("in", "desp_poupanca", "Despesas de depósitos de poupança"),
    ("in", "desp_rdc", "Despesas de depósitos a prazo (RDC)"),
    ("in", "desp_lca", "Despesas de LCA"),
    ("in", "desp_lfs", "Despesas de LFS"),
    ("in", "fgcoop_neg", "Contribuição ao FGCoop"),
    ("sub", "desp_captacao", "(=) Operações de captação no mercado", ["desp_poupanca", "desp_rdc", "desp_lca", "desp_lfs", "fgcoop_neg"]),
    ("in", "desp_repasses_interfin", "Repasses interfinanceiros (Central e Banco)"),
    ("in", "desp_emprestimos_repasses", "Empréstimos e repasses (BNDES e outros)"),
    ("sub", "desp_repasses", "(=) Operações de empréstimos e repasses", ["desp_repasses_interfin", "desp_emprestimos_repasses"]),
    ("in", "pdd", "Provisão para perdas esperadas associadas ao risco de crédito"),
    ("sub", "despesas_intermed", "(=) Total das despesas da intermediação", ["desp_captacao", "desp_repasses", "pdd"]),
    ("tot", "resultado_bruto", "RESULTADO BRUTO DA INTERMEDIAÇÃO FINANCEIRA", ["receitas_intermed", "despesas_intermed"]),
    ("sec", None, "OUTRAS RECEITAS E DESPESAS OPERACIONAIS"),
    ("in", "servicos", "Receitas de prestação de serviços"),
    ("in", "tarifas", "Rendas de tarifas"),
    ("in", "pessoal", "Despesas de pessoal"),
    ("in", "outras_adm", "Outras despesas administrativas"),
    ("in", "tributarias", "Despesas tributárias"),
    ("in", "outras_rec_op", "Outras receitas operacionais"),
    ("in", "outras_desp_op", "Outras despesas operacionais"),
    ("in", "contingencias", "Provisões para contingências e garantias (líquidas)"),
    ("sub", "outras_op", "(=) Total de outras receitas e despesas operacionais",
     ["servicos", "tarifas", "pessoal", "outras_adm", "tributarias", "outras_rec_op", "outras_desp_op", "contingencias"]),
    ("tot", "resultado_operacional", "RESULTADO OPERACIONAL", ["resultado_bruto", "outras_op"]),
    ("in", "nao_operacional", "Resultado não operacional"),
    ("tot", "resultado_antes_ir", "RESULTADO ANTES DA TRIBUTAÇÃO", ["resultado_operacional", "nao_operacional"]),
    ("in", "ir_cs", "IR e CS sobre atos não cooperativos"),
    ("tot", "resultado", "RESULTADO ANTES DOS JUROS AO CAPITAL (sobras brutas)", ["resultado_antes_ir", "ir_cs"]),
    ("in", "juros_capital", "Juros ao capital"),
    ("tot", "sobras_apos_juros", "SOBRAS DO EXERCÍCIO APÓS JUROS AO CAPITAL", ["resultado", "juros_capital"]),
    ("sec", None, "DESTINAÇÕES"),
    ("in", "dest_fates_ato", "FATES: resultado de atos não cooperativos"),
    ("fx", "dest_reserva", "Reserva legal (45% da base)", None),
    ("fx", "dest_fates", "FATES estatutário (5% da base)", None),
    ("tot", "sobras_ago", "SOBRAS À DISPOSIÇÃO DA AGO", ["sobras_apos_juros", "dest_fates_ato", "dest_reserva", "dest_fates"]),
]

ws = wb.create_sheet("DSP")
titulo(ws, "Demonstração de Sobras ou Perdas (DSP): Coop Vale do Ipê", "Em R$ mil | exercícios de 2024 e 2025 | AV sobre o total das receitas da intermediação | despesas com sinal negativo")
cabecalho(ws, 4, ["Linha", "2024", "2025", "AV 2024", "AV 2025", "AH 25/24", "Variação (R$ mil)", "O que entra"])
_r = 5
for ln in LINHAS_DSP:
    if ln[0] != "sec":
        REG[("DSP", ln[1])] = _r
    _r += 1
PCT_RL_LINHA = None


def _fx_reserva(col):
    b = f"({col}{REG[('DSP', 'sobras_apos_juros')]}+{col}{REG[('DSP', 'dest_fates_ato')]})"
    return f"=-ROUND({b}*{par('pct_reserva')},-1)"


def _fx_fates(col):
    b = f"({col}{REG[('DSP', 'sobras_apos_juros')]}+{col}{REG[('DSP', 'dest_fates_ato')]})"
    return f"=-ROUND({b}*{par('pct_fates')},-1)"


def valor_dsp(k, a):
    if k == "fgcoop_neg":
        return -D[a]["fgcoop"]
    return D[a][k]


def base_dsp(k):
    b = REG[("DSP", "receitas_intermed")]
    return (b, b)


# Dados precisa existir antes por causa de par(); registramos as linhas de parâmetros já aqui
LINHAS_DSP = [ln if ln[1] not in ("dest_reserva", "dest_fates") else
              (ln[0], ln[1], ln[2], _fx_reserva if ln[1] == "dest_reserva" else _fx_fates) for ln in LINHAS_DSP]

# ================================================================== Dados (registro antecipado)
wsD = wb.create_sheet("Dados")
DADOS = []  # (tipo, chave, rótulo, valores{ano: v} ou valor, fmt, nota)


def dsec(t):
    DADOS.append(("sec", None, t, None, None, None))


def dano(k, rot, vals, fmt=F_MIL, nota=""):
    DADOS.append(("ano", k, rot, vals, fmt, nota))


def dpar(k, rot, v, fmt=F_PCT1, nota=""):
    DADOS.append(("par", k, rot, v, fmt, nota))


dsec("PREMISSAS MACRO")
dano("cdi", "CDI médio do ano", {2024: M.CDI[2024], 2025: M.CDI[2025]}, F_PCT2, "Premissa do exemplo (aproximada)")
dano("ipca", "IPCA do ano", {2024: M.IPCA[2024], 2025: M.IPCA[2025]}, F_PCT2, "Premissa do exemplo")
dsec("SALDOS PARA MÉDIAS (2023 digitado; 2024 e 2025 ligados ao BP)")
SALDO_KEYS = [
    ("ativo", "Ativo total"), ("caixa", "Disponibilidades"), ("centralizacao", "Centralização financeira"),
    ("aplic_interfin", "Aplicações interfinanceiras"), ("tvm", "TVM"), ("liq_rentavel", "Liquidez rentável (centralização + aplicações + TVM)"),
    ("liquidez", "Liquidez total (caixa + liquidez rentável)"), ("cred_comercial", "Carteira comercial (empréstimos + financiamentos)"),
    ("cred_rural", "Carteira rural"), ("carteira", "Carteira bruta"), ("provisao", "Provisão (negativa)"),
    ("dep_vista", "Depósitos à vista"), ("poupanca", "Poupança"), ("rdc", "RDC"), ("depositos", "Depósitos totais"),
    ("lca", "LCA"), ("captacao", "Captação com associados (depósitos + LCA)"), ("repasses_interfin", "Repasses interfinanceiros"),
    ("emprestimos_repasses", "Empréstimos e repasses"), ("lfs", "LFS"), ("pl", "Patrimônio líquido"), ("capital_social", "Capital social"),
]
for k, rot in SALDO_KEYS:
    dano("s_" + k, rot, None, F_MIL)
dsec("CARTEIRA POR FAIXA DE ATRASO (saldo bruto)")
for fk, rot in (("em_dia", "Em dia"), ("d1_14", "1 a 14 dias"), ("d15_30", "15 a 30 dias"), ("d31_60", "31 a 60 dias"),
                ("d61_90", "61 a 90 dias"), ("d90", "Acima de 90 dias")):
    dano("at_" + fk, rot, {a: NOTAS["atraso"][a][fk] for a in M.ANOS})
dano("at_total", "Total (confere com a carteira do BP)", None)
dsec("CARTEIRA POR ESTÁGIO (Res. CMN 4.966)")
for e, rot in (("e1", "Estágio 1: saldo"), ("e2", "Estágio 2: saldo"), ("e3", "Estágio 3: saldo")):
    dano("est_" + e, rot, {a: NOTAS["estagio"][a][e][0] for a in M.ANOS})
for e, rot in (("e1", "Estágio 1: provisão"), ("e2", "Estágio 2: provisão"), ("e3", "Estágio 3: provisão")):
    dano("prov_" + e, rot, {a: NOTAS["estagio"][a][e][1] for a in M.ANOS})
dano("prov_comercial", "Provisão da carteira comercial", {2024: NOTAS["prov_comercial"][2024], 2025: NOTAS["prov_comercial"][2025]})
dano("prov_rural", "Provisão da carteira rural", {2024: NOTAS["prov_rural"][2024], 2025: NOTAS["prov_rural"][2025]})
dano("renegociadas", "Operações renegociadas no ano (saldo)", {a: NOTAS["renegociadas"][a] for a in M.ANOS})
dano("top20_devedores", "20 maiores devedores ÷ carteira", {a: NOTAS["top20_devedores"][a] for a in M.ANOS}, F_PCT1)
dsec("RECEITA DE CRÉDITO POR CARTEIRA (detalhe da DSP)")
dano("rend_comercial", "Receita da carteira comercial", {a: D[a]["rend_comercial"] for a in M.ANOS})
dano("rend_rural", "Receita da carteira rural", {a: D[a]["rend_rural"] for a in M.ANOS})
dsec("ASSOCIADOS E CONCENTRAÇÃO")
dano("associados", "Associados (quantidade)", {a: NOTAS["associados"][a] for a in (2023, 2024, 2025)}, '#,##0')
dano("top10_depositantes", "10 maiores depositantes ÷ captação", {a: NOTAS["top10_depositantes"][a] for a in M.ANOS}, F_PCT1)
dsec("ITENS NÃO RECORRENTES (efeito no resultado: + ganho, - perda)")
dano("nr_pdv", "Programa de desligamento voluntário (em despesas de pessoal)", {a: NOTAS["nr_pdv"][a] for a in M.ANOS})
dano("nr_venda_imovel", "Venda de imóvel de agência (resultado não operacional)", {a: NOTAS["nr_venda_imovel"][a] for a in M.ANOS})
dano("nr_reversao_conting", "Reversão de contingência fiscal ganha (em contingências)", {a: NOTAS["nr_reversao_conting"][a] for a in M.ANOS})
dano("nr_recuperacao_extra", "Recuperação extraordinária de um grande crédito baixado", {a: NOTAS["nr_recuperacao_extra"][a] for a in M.ANOS})
dano("nr_total", "Total de itens não recorrentes", None)
dsec("DETALHE DE DESPESAS ADMINISTRATIVAS")
dano("deprec", "Depreciação e amortização (dentro de outras despesas administrativas)", {a: NOTAS["deprec"][a] for a in M.ANOS})
dano("rateios", "Rateios da Central e da Confederação (dentro de outras despesas administrativas)", {a: NOTAS["rateios"][a] for a in M.ANOS})
dsec("CAPITAL REGULATÓRIO (premissas)")
dano("rwa_mercado", "RWA de risco de mercado (informado)", {a: NOTAS["rwa_mercado"][a] for a in M.ANOS})
dano("ajuste_prud_outros", "Outros ajustes prudenciais (além do intangível)", {a: NOTAS["ajuste_prudencial_outros"][a] for a in M.ANOS})
dano("receita_bruta", "Receita bruta para risco operacional (2022 em B? ver nota)", None, F_MIL, "2024 e 2025 = margem financeira + serviços e tarifas (aba Gerencial)")
dpar("receita_bruta_2022", "Receita bruta de 2022", NOTAS["receita_bruta_2022"], F_MIL)
dpar("receita_bruta_2023", "Receita bruta de 2023", NOTAS["receita_bruta_2023"], F_MIL)
dpar("fator_op", "Fator do risco operacional (abordagem simplificada)", 0.15, F_PCT1, "Didático: 15% da receita bruta média")
dpar("fator_f", "Fator F (exigência básica de PR)", 0.08, F_PCT1)
dpar("exig_total", "Exigência de PR + adicional de conservação (S4)", 0.105, F_PCT1, "8% + 2,5%")
dpar("piso_gerencial", "Piso gerencial de MM/RWA (política de recursos corporativos)", 0.02, F_PCT1, "Aula de Recursos Corporativos: 2 p.p.")
dpar("fpr_politica", "FPR usado na conta do potencial de crédito (política)", 1.10, F_PCT1)
dpar("nml", "Nível mínimo de liquidez (NML) sobre depósitos", NOTAS["nml"], F_PCT1, "Ilustrativo")
dpar("pct_reserva", "Destinação à reserva legal (estatuto)", M.PCT_RESERVA_LEGAL, F_PCT1, "Lei 5.764: mínimo de 10%")
dpar("pct_fates", "Destinação ao FATES (estatuto)", M.PCT_FATES, F_PCT1, "Lei 5.764: mínimo de 5%")
dsec("FATORES DE PONDERAÇÃO DE RISCO (FPR, ilustrativos)")
FPR_ROT = {"caixa": "Disponibilidades", "centralizacao": "Centralização financeira", "aplic_interfin": "Aplicações interfinanceiras",
           "tvm": "Títulos públicos federais", "comercial": "Carteira comercial (líquida de provisão)", "rural": "Carteira rural (líquida de provisão)",
           "outros_ativos_fin": "Outros ativos financeiros", "ativos_fiscais": "Ativos fiscais", "bndu_outros": "BNDU e outros ativos",
           "investimentos": "Investimentos", "imobilizado": "Imobilizado", "coobrigacoes": "Coobrigações e garantias (fora do balanço)"}
for k, rot in FPR_ROT.items():
    dpar("fpr_" + k, "FPR: " + rot, M.FPR[k], '0%')
dsec("FTP: TAXAS DE TRANSFERÊNCIA (anuais)")
FTP_ROT = {"comercial": "Crédito comercial (paga à tesouraria)", "rural": "Crédito rural (paga à tesouraria)",
           "dep_vista": "Depósito à vista, parcela estável (recebe da tesouraria)", "poupanca": "Poupança (recebe)",
           "rdc": "RDC (recebe)", "lca": "LCA (recebe)"}
for k, rot in FTP_ROT.items():
    dano("ftp_" + k, "FTP: " + rot, {a: M.FTP[k][a] for a in M.ANOS}, F_PCT2)
dsec("RAROC POR CARTEIRA (premissas de 2025)")
dpar("raroc_custo_comercial", "Custo operacional alocado: comercial (% do saldo médio)", M.RAROC["custo_op"]["comercial"], F_PCT2)
dpar("raroc_custo_rural", "Custo operacional alocado: rural", M.RAROC["custo_op"]["rural"], F_PCT2)
dpar("raroc_pe_comercial", "Perda esperada: comercial (% do saldo médio)", M.RAROC["pe"]["comercial"], F_PCT2)
dpar("raroc_pe_rural", "Perda esperada: rural", M.RAROC["pe"]["rural"], F_PCT2)
dpar("raroc_capital", "Capital alocado (% do RWA)", M.RAROC["capital"], F_PCT2, "10,5% + 2 p.p.")
dpar("raroc_hurdle", "Retorno mínimo exigido (hurdle)", M.RAROC["hurdle"], F_PCT1, "Custo de capital da aula de Viabilidade")
dsec("LCR DIDÁTICO: PREMISSAS DE ESTRESSE DE 30 DIAS")
LCR_ROT = [("dv_fator", "Fuga de depósitos à vista", F_PCT1), ("poup_fator", "Fuga de poupança", F_PCT1),
           ("rdc_fator", "Fuga de RDC (média, inclui liquidez diária)", F_PCT1), ("lca_30d", "LCA vencendo em 30 dias (sai 100%)", F_MIL),
           ("obrig_30d", "Repasses e obrigações vencendo em 30 dias (sai 100%)", F_MIL),
           ("limites", "Limites de crédito não utilizados (cheque especial, cartão)", F_MIL),
           ("limites_fator", "Uso dos limites no estresse", F_PCT1), ("outras_saidas", "Outras saídas contratuais", F_MIL),
           ("parcelas_30d", "Parcelas de crédito a receber em 30 dias", F_MIL), ("parcelas_fator", "Parcela das entradas de crédito considerada", F_PCT1),
           ("aplic_30d", "Aplicações interfinanceiras vencendo em 30 dias (entram 100%)", F_MIL)]
for k, rot, fmt in LCR_ROT:
    dano("lcr_" + k, rot, {a: M.LCR[a][k] for a in M.ANOS}, fmt)
dpar("lcr_teto_entradas", "Teto das entradas (% das saídas)", 0.75, F_PCT1)
dsec("FLUXOS DO PL EM 2025 (para a DMPL)")
dpar("dmpl_integralizacoes", "Integralizações de capital (novos associados e aportes)", M.DMPL["integralizacoes"], F_MIL)
dpar("dmpl_devolucoes", "Devoluções de quotas (associados desligados)", M.DMPL["devolucoes"], F_MIL)
dpar("dmpl_pct_ago_capital", "AGO 2025: parte das sobras de 2024 capitalizada", M.DMPL["pct_ago_capital"], F_PCT1, "O restante foi creditado em conta corrente")

r = 4
for tipo, k, rot, v, fmt, nota in DADOS:
    if tipo == "sec":
        r += 1 if r > 4 else 0
    if tipo != "sec":
        REG[("Dados", k)] = r
    r += 1


def link_saldo(k, a):
    if k in ("liq_rentavel",):
        return f"={da('s_centralizacao', a)}+{da('s_aplic_interfin', a)}+{da('s_tvm', a)}"
    if k == "liquidez":
        return f"={da('s_caixa', a)}+{da('s_liq_rentavel', a)}"
    if k == "captacao":
        return f"={da('s_depositos', a)}+{da('s_lca', a)}"
    if a == 2023:
        if k == "carteira":
            return f"={da('s_cred_comercial', a)}+{da('s_cred_rural', a)}"
        if k == "depositos":
            return f"={da('s_dep_vista', a)}+{da('s_poupanca', a)}+{da('s_rdc', a)}"
        return S[k][2023]
    if k == "cred_comercial":
        return f"={bp('cred_emprestimos', a)}+{bp('cred_financiamentos', a)}"
    return f"={bp(k, a)}"


cabecalho(wsD, 3, ["Item", "2023 / valor", "2024", "2025", "Nota"])
titulo(wsD, "Notas explicativas e premissas da Coop Vale do Ipê", "Em R$ mil, salvo indicação. Azul = dado digitado; verde = link; preto = fórmula.", 5)
r = 4
for tipo, k, rot, v, fmt, nota in DADOS:
    if tipo == "sec":
        if r > 4:
            r += 1
        c = wsD.cell(row=r, column=1, value=rot)
        for cc in range(1, 6):
            estilo(wsD.cell(row=r, column=cc), COR_TITULO, True, fundo=COR_SECAO)
        r += 1
        continue
    assert REG[("Dados", k)] == r, (k, r)
    wsD.cell(row=r, column=1, value=rot)
    estilo(wsD.cell(row=r, column=1))
    if tipo == "par":
        c = wsD.cell(row=r, column=2, value=v)
        estilo(c, AZUL, fmt=fmt)
    else:
        for a, col in COL_ANO["Dados"].items():
            if k.startswith("s_"):
                val = link_saldo(k[2:], a)
            elif k == "at_total":
                if a == 2023:
                    continue
                val = f"=SUM({col}{REG[('Dados', 'at_em_dia')]}:{col}{REG[('Dados', 'at_d90')]})"
            elif k == "nr_total":
                if a == 2023:
                    continue
                val = f"=SUM({col}{REG[('Dados', 'nr_pdv')]}:{col}{REG[('Dados', 'nr_recuperacao_extra')]})"
            elif k == "receita_bruta":
                if a == 2023:
                    continue
                val = None  # preenchido depois que a aba Gerencial existir
            else:
                if a not in v:
                    continue
                val = v[a]
            if val is None:
                continue
            c = wsD.cell(row=r, column="ABCDE".index(col) + 1, value=val)
            estilo(c, cor_de(val), fmt=fmt)
    if nota:
        estilo(wsD.cell(row=r, column=5, value=nota), "595959", italico=True, tam=9)
    r += 1
larguras(wsD, [70, 14, 14, 14, 60])
wsD.freeze_panes = "B4"

# agora a DSP (usa par())
ws = wb["DSP"]
fim_dsp = montar_demonstrativo(ws, LINHAS_DSP, base_dsp, 5, "DSP", valor_dsp, DESC_DSP)
larguras(ws, [62, 14, 14, 9, 9, 10, 14, 90])
ws.freeze_panes = "B5"
ws.cell(row=fim_dsp + 1, column=1, value="Conferência: sobras à disposição da AGO = sobras acumuladas do BP")
estilo(ws.cell(row=fim_dsp + 1, column=1), "595959", italico=True)
for a, col in COL_ANO["DSP"].items():
    v = f"={col}{REG[('DSP', 'sobras_ago')]}-{bp('sobras_acum', a)}"
    estilo(ws.cell(row=fim_dsp + 1, column="ABC".index(col) + 1, value=v), PRETO, fmt=F_MIL)

# ================================================================== Gerencial
ws = wb.create_sheet("Gerencial")
titulo(ws, "DRE gerencial em cascata (a mesma DSP, arrumada para análise)", "Em R$ mil | é a visão do simulador de resultado: margem financeira, receita líquida, custo e risco", 5)
cabecalho(ws, 4, ["Linha", "2024", "2025", "AH 25/24", "Como se calcula"])
GER = [
    ("g_credito", "Receita de crédito", lambda a: f"={ds('rend_credito', a)}", "Operações de crédito da DSP"),
    ("g_liquidez", "(+) Receita da liquidez", lambda a: f"={ds('rend_centralizacao', a)}+{ds('rend_aplic', a)}+{ds('rend_tvm', a)}", "Centralização + aplicações interfinanceiras + TVM"),
    ("g_captacao", "(-) Despesa de captação", lambda a: f"={ds('desp_captacao', a)}", "Poupança, RDC, LCA, LFS e FGCoop"),
    ("g_repasses", "(-) Despesa de repasses", lambda a: f"={ds('desp_repasses', a)}", "Central, Banco e BNDES"),
    ("mf", "(=) MARGEM FINANCEIRA", lambda a: f"=SUM(B{{g_credito}}:B{{g_repasses}})", "Antes da PDD. É o 'spread' da cooperativa inteira"),
    ("g_servicos", "(+) Serviços e tarifas", lambda a: f"={ds('servicos', a)}+{ds('tarifas', a)}", ""),
    ("g_outras", "(+/-) Outras receitas e despesas operacionais", lambda a: f"={ds('outras_rec_op', a)}+{ds('outras_desp_op', a)}+{ds('tributarias', a)}+{ds('contingencias', a)}", "Outras receitas - outras despesas - tributárias +/- contingências"),
    ("receita_liquida", "(=) RECEITA LÍQUIDA (base do índice de eficiência)", None, "Margem + serviços + outras"),
    ("desp_adm", "(-) Despesas administrativas", lambda a: f"={ds('pessoal', a)}+{ds('outras_adm', a)}", "Pessoal + outras administrativas"),
    ("pdd_liquida", "(-) PDD líquida de recuperações", lambda a: f"={ds('pdd', a)}+{ds('recuperacao', a)}", "Custo do risco de crédito do ano"),
    ("res_op", "(=) RESULTADO OPERACIONAL", None, ""),
    ("g_nao_op", "(+/-) Resultado não operacional", lambda a: f"={ds('nao_operacional', a)}", ""),
    ("g_ir", "(-) IR e CS", lambda a: f"={ds('ir_cs', a)}", ""),
    ("resultado", "(=) RESULTADO ANTES DOS JUROS AO CAPITAL", None, "Base do ROE"),
    ("check", "Conferência com a DSP (deve ser zero)", lambda a: f"=B{{resultado}}-{ds('resultado', a)}", ""),
    ("nr", "(-) Itens não recorrentes (notas)", lambda a: f"={da('nr_total', a)}", "Ganhos e perdas que não se repetem"),
    ("resultado_recorrente", "(=) RESULTADO RECORRENTE", None, "Resultado que tende a se repetir"),
]
for n, g in enumerate(GER):
    REG[("Gerencial", g[0])] = 5 + n
for n, (k, rot, fn, expl) in enumerate(GER):
    r = 5 + n
    ws.cell(row=r, column=1, value=rot)
    tot = rot.startswith("(=)")
    for a, col in COL_ANO["Gerencial"].items():
        if k == "receita_liquida":
            v = f"={col}{REG[('Gerencial', 'mf')]}+{col}{REG[('Gerencial', 'g_servicos')]}+{col}{REG[('Gerencial', 'g_outras')]}"
        elif k == "res_op":
            v = f"={col}{REG[('Gerencial', 'receita_liquida')]}+{col}{REG[('Gerencial', 'desp_adm')]}+{col}{REG[('Gerencial', 'pdd_liquida')]}"
        elif k == "resultado":
            v = f"={col}{REG[('Gerencial', 'res_op')]}+{col}{REG[('Gerencial', 'g_nao_op')]}+{col}{REG[('Gerencial', 'g_ir')]}"
        elif k == "resultado_recorrente":
            v = f"={col}{REG[('Gerencial', 'resultado')]}-{col}{REG[('Gerencial', 'nr')]}"
        else:
            v = fn(a)
            v = v.replace("B{g_credito}", f"{col}{REG[('Gerencial', 'g_credito')]}").replace("B{g_repasses}", f"{col}{REG[('Gerencial', 'g_repasses')]}")
            v = v.replace("B{resultado}", f"{col}{REG[('Gerencial', 'resultado')]}")
        c = ws.cell(row=r, column="ABC".index(col) + 1, value=v)
        estilo(c, cor_de(v), tot, F_MIL, COR_SUB if tot else None)
    estilo(ws.cell(row=r, column=1), PRETO, tot, fundo=COR_SUB if tot else None)
    if k != "check":
        estilo(ws.cell(row=r, column=4, value=f'=IFERROR(C{r}/B{r}-1,"n.a.")'), PRETO, tot, F_PCT1, COR_SUB if tot else None)
    estilo(ws.cell(row=r, column=5, value=expl), "404040", tam=9)
larguras(ws, [52, 14, 14, 10, 60])

# receita bruta (Dados) agora que Gerencial existe
rr = REG[("Dados", "receita_bruta")]
wsD.cell(row=rr, column=1, value="Receita bruta para risco operacional (margem financeira + serviços e tarifas)")
for a, col in ((2024, "C"), (2025, "D")):
    v = f"={ge('mf', a)}+{ge('g_servicos', a)}"
    estilo(wsD.cell(row=rr, column="ABCD".index(col) + 1, value=v), PRETO, fmt=F_MIL)

# ================================================================== DMPL
ws = wb.create_sheet("DMPL")
titulo(ws, "Demonstração das mutações do PL em 2025 (DMPL)", "Em R$ mil | cada linha mostra o efeito em cada conta do PL", 7)
cabecalho(ws, 4, ["Evento", "Capital social", "Reserva legal", "Sobras acumuladas", "Ajustes de avaliação", "Total do PL", "Explicação"])
sob24 = bp("sobras_acum", 2024)
DM = [
    ("ini", "Saldo em 31/12/2024", [f"={bp('capital_social', 2024)}", f"={bp('reserva_legal', 2024)}", f"={bp('sobras_acum', 2024)}", f"={bp('aap', 2024)}"], "Vem do BP de 2024"),
    ("integ", "Integralizações de capital", [f"={par('dmpl_integralizacoes')}", 0, 0, 0], "Novos associados e aportes"),
    ("devol", "Devoluções de quotas", [f"={par('dmpl_devolucoes')}", 0, 0, 0], "Associados que saíram levam a quota"),
    ("ago_cap", "AGO: capitalização de sobras de 2024", [f"=ROUND({sob24}*{par('dmpl_pct_ago_capital')},-1)", 0, f"=-B{{ago_cap}}", 0], "Sai das sobras, entra no capital: PL não muda"),
    ("ago_dist", "AGO: sobras de 2024 creditadas em conta corrente", [0, 0, f"=-({sob24}+D{{ago_cap}})", 0], "Sai do PL e vira depósito à vista"),
    ("res", "Resultado do exercício (antes dos juros ao capital)", [0, 0, f"={ds('resultado', 2025)}", 0], "DSP de 2025"),
    ("jcp", "Juros ao capital: destinação", [0, 0, f"={ds('juros_capital', 2025)}", 0], "Remuneração das quotas"),
    ("jcp_inc", "Juros ao capital: incorporação às quotas", [f"=-{ds('juros_capital', 2025)}", 0, 0, 0], "Creditados no capital do associado: efeito zero no PL"),
    ("fates", "FATES (ato não cooperativo + estatutário)", [0, 0, f"={ds('dest_fates_ato', 2025)}+{ds('dest_fates', 2025)}", 0], "Vai para o passivo (obrigações sociais)"),
    ("reserva", "Reserva legal", [0, f"=-{ds('dest_reserva', 2025)}", f"={ds('dest_reserva', 2025)}", 0], "Sai das sobras, entra na reserva"),
    ("aap", "Ajuste de avaliação patrimonial do ano", [0, 0, 0, f"={bp('aap', 2025)}-{bp('aap', 2024)}"], "Marcação a mercado (VJORA)"),
    ("fim", "Saldo em 31/12/2025", None, "Soma das linhas"),
    ("chk", "Conferência com o BP de 2025 (deve ser zero)", None, ""),
]
for n, d in enumerate(DM):
    REG[("DMPL", d[0])] = 5 + n
for n, (k, rot, vals, expl) in enumerate(DM):
    r = 5 + n
    ws.cell(row=r, column=1, value=rot)
    tot = k in ("ini", "fim")
    estilo(ws.cell(row=r, column=1), PRETO, tot, fundo=COR_SUB if tot else None)
    for j, col in enumerate("BCDE"):
        if k == "fim":
            v = f"=SUM({col}{REG[('DMPL', 'ini')]}:{col}{REG[('DMPL', 'aap')]})"
        elif k == "chk":
            alvo = {"B": "capital_social", "C": "reserva_legal", "D": "sobras_acum", "E": "aap"}[col]
            v = f"={col}{REG[('DMPL', 'fim')]}-{bp(alvo, 2025)}"
        else:
            v = vals[j]
            if isinstance(v, str):
                v = v.replace("B{ago_cap}", f"B{REG[('DMPL', 'ago_cap')]}").replace("D{ago_cap}", f"D{REG[('DMPL', 'ago_cap')]}")
        c = ws.cell(row=r, column=j + 2, value=v)
        estilo(c, cor_de(v), tot, F_MIL, COR_SUB if tot else None)
    c = ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})")
    estilo(c, PRETO, True, F_MIL, COR_SUB if tot else None)
    estilo(ws.cell(row=r, column=7, value=expl), "404040", tam=9)
larguras(ws, [50, 14, 14, 16, 14, 14, 55])

# ================================================================== RWA
ws = wb.create_sheet("RWA")
titulo(ws, "RWA, Patrimônio de Referência e Índice de Basileia", "Em R$ mil | FPR ilustrativos (aba Dados) | exigência do S4: 10,5% do RWA", 6)
cabecalho(ws, 4, ["Exposição", "FPR", "Exposição 2024", "RWA 2024", "Exposição 2025", "RWA 2025"])
EXPO = [
    ("caixa", lambda a: f"={bp('caixa', a)}"),
    ("centralizacao", lambda a: f"={bp('centralizacao', a)}"),
    ("aplic_interfin", lambda a: f"={bp('aplic_interfin', a)}"),
    ("tvm", lambda a: f"={bp('tvm', a)}"),
    ("comercial", lambda a: f"={bp('cred_emprestimos', a)}+{bp('cred_financiamentos', a)}-{da('prov_comercial', a)}"),
    ("rural", lambda a: f"={bp('cred_rural', a)}-{da('prov_rural', a)}"),
    ("outros_ativos_fin", lambda a: f"={bp('outros_ativos_fin', a)}"),
    ("ativos_fiscais", lambda a: f"={bp('ativos_fiscais', a)}"),
    ("bndu_outros", lambda a: f"={bp('bndu', a)}+{bp('outros_ativos', a)}"),
    ("investimentos", lambda a: f"={bp('investimentos', a)}"),
    ("imobilizado", lambda a: f"={bp('imobilizado', a)}"),
    ("coobrigacoes", lambda a: f"={bp('coobrigacoes', a)}"),
]
r = 5
for k, fn in EXPO:
    REG[("RWA", k)] = r
    ws.cell(row=r, column=1, value=FPR_ROT[k])
    estilo(ws.cell(row=r, column=1))
    estilo(ws.cell(row=r, column=2, value=f"={par('fpr_' + k)}"), VERDE, fmt="0%")
    for a, ce, cr in ((2024, "C", "D"), (2025, "E", "F")):
        v = fn(a)
        estilo(ws.cell(row=r, column="ABCDEF".index(ce) + 1, value=v), cor_de(v), fmt=F_MIL)
        estilo(ws.cell(row=r, column="ABCDEF".index(cr) + 1, value=f"={ce}{r}*$B{r}"), PRETO, fmt=F_MIL)
    r += 1
r_ini, r_fim = 5, r - 1
LIN_RWA = [
    ("rwa_credito", "RWA de risco de crédito (soma)", lambda c, a: f"=SUM({c}{r_ini}:{c}{r_fim})", F_MIL, True),
    ("rb_media", "Receita bruta média de 3 anos", lambda c, a: (f"=AVERAGE({par('receita_bruta_2022')},{par('receita_bruta_2023')},{da('receita_bruta', 2024)})" if a == 2024
                                                               else f"=AVERAGE({par('receita_bruta_2023')},{da('receita_bruta', 2024)},{da('receita_bruta', 2025)})"), F_MIL, False),
    ("rwa_operacional", "RWA de risco operacional = fator × receita média ÷ F", lambda c, a: f"={par('fator_op')}*{c}{{rb_media}}/{par('fator_f')}", F_MIL, True),
    ("rwa_mercado", "RWA de risco de mercado", lambda c, a: f"={da('rwa_mercado', a)}", F_MIL, False),
    ("rwa", "RWA TOTAL", lambda c, a: f"={c}{{rwa_credito}}+{c}{{rwa_operacional}}+{c}{{rwa_mercado}}", F_MIL, True),
    ("densidade", "Densidade: RWA ÷ ativo total", lambda c, a: f"={c}{{rwa}}/{bp('ativo', a)}", F_PCT1, False),
    ("pl", "Patrimônio líquido", lambda c, a: f"={bp('pl', a)}", F_MIL, False),
    ("ajustes", "(-) Ajustes prudenciais (intangível + outros)", lambda c, a: f"=-({bp('intangivel', a)}+{da('ajuste_prud_outros', a)})", F_MIL, False),
    ("cp", "(=) Capital principal (Nível I)", lambda c, a: f"={c}{{pl}}+{c}{{ajustes}}", F_MIL, True),
    ("n2", "(+) Nível II: LFS", lambda c, a: f"={bp('lfs', a)}", F_MIL, False),
    ("pr", "(=) PATRIMÔNIO DE REFERÊNCIA (PR)", lambda c, a: f"={c}{{cp}}+{c}{{n2}}", F_MIL, True),
    ("basileia", "Índice de Basileia = PR ÷ RWA", lambda c, a: f"={c}{{pr}}/{c}{{rwa}}", F_PCT2, True),
    ("indice_cp", "Índice de capital principal = CP ÷ RWA", lambda c, a: f"={c}{{cp}}/{c}{{rwa}}", F_PCT2, False),
    ("mm_rwa", "MM/RWA = (Basileia - 10,5%) em p.p.", lambda c, a: f"=({c}{{basileia}}-{par('exig_total')})*100", F_PP, True),
    ("margem_capital", "Margem de capital (R$ mil) = PR - 10,5% × RWA", lambda c, a: f"={c}{{pr}}-{par('exig_total')}*{c}{{rwa}}", F_MIL, False),
    ("potencial", "Potencial de crédito até MM/RWA de 2 p.p. = (PR ÷ 12,5% - RWA) ÷ 1,1", lambda c, a: f"=({c}{{pr}}/({par('exig_total')}+{par('piso_gerencial')})-{c}{{rwa}})/{par('fpr_politica')}", F_MIL, False),
    ("basileia_sem_lfs", "Índice de Basileia sem a LFS", lambda c, a: f"={c}{{cp}}/{c}{{rwa}}", F_PCT2, False),
]
r += 1
for n, it in enumerate(LIN_RWA):
    REG[("RWA", it[0])] = r + n
for n, (k, rot, fn, fmt, neg) in enumerate(LIN_RWA):
    rr = r + n
    ws.cell(row=rr, column=1, value=rot)
    estilo(ws.cell(row=rr, column=1), PRETO, neg, fundo=COR_SUB if neg else None)
    for a, c in ((2024, "D"), (2025, "F")):
        v = fn(c, a)
        for kk, _r2 in [(x[0], REG[("RWA", x[0])]) for x in LIN_RWA]:
            v = v.replace("{" + kk + "}", str(_r2))
        estilo(ws.cell(row=rr, column="ABCDEF".index(c) + 1, value=v), cor_de(v), neg, fmt, COR_SUB if neg else None)
larguras(ws, [62, 8, 15, 15, 15, 15])

# ================================================================== Indicadores
ws = wb.create_sheet("Indicadores")
titulo(ws, "Os 15 indicadores calculados (2024 e 2025)", "Cada linha é uma fórmula viva. Coluna E: variação (p.p. para taxas, % para valores).", 6)
cabecalho(ws, 4, ["Indicador", "Como se calcula", "2024", "2025", "Variação", "O que olhar"])


def rec(k, a):
    return ge(k, a)


IND = []


def fam(t):
    IND.append(("sec", t))


def ind(k, rot, como, fn, fmt, leitura="", var="pp"):
    IND.append(("ind", k, rot, como, fn, fmt, leitura, var))


fam("1. PATRIMÔNIO LÍQUIDO (PL)")
ind("pl", "Patrimônio líquido (R$ mil)", "Capital social + reservas + sobras + ajustes", lambda a: f"={bp('pl', a)}", F_MIL, "Colchão para perdas e base para crescer", "ah")
ind("cresc_pl", "Crescimento do PL", "PL do ano ÷ PL do ano anterior - 1", lambda a: f"={da('s_pl', a)}/{da('s_pl', a - 1)}-1", F_PCT1, "Compare com o crescimento do RWA")
ind("pl_ativo", "PL ÷ ativo total", "PL ÷ ativo", lambda a: f"={bp('pl', a)}/{bp('ativo', a)}", F_PCT1, "Alavancagem contábil (sem risco)")
ind("alavancagem", "Ativo ÷ PL (vezes)", "Ativo ÷ PL", lambda a: f"={bp('ativo', a)}/{bp('pl', a)}", F_X, "Quantos reais de ativo por real de capital", "x")
ind("cap_pl", "Capital social ÷ PL", "Quotas ÷ PL", lambda a: f"={bp('capital_social', a)}/{bp('pl', a)}", F_PCT1, "Quanto do PL pode sair com resgate de quotas")
ind("cap_assoc", "Capital social por associado (R$)", "Capital × 1.000 ÷ associados", lambda a: f"={bp('capital_social', a)}*1000/{da('associados', a)}", F_RS, "", "ah")
fam("2. CRESCIMENTO DA CARTEIRA DE CRÉDITO")
ind("carteira", "Carteira bruta (R$ mil)", "Empréstimos + financiamentos + rural", lambda a: f"={bp('carteira', a)}", F_MIL, "", "ah")
ind("cresc_carteira", "Crescimento nominal da carteira", "Carteira ÷ carteira anterior - 1", lambda a: f"={da('s_carteira', a)}/{da('s_carteira', a - 1)}-1", F_PCT1, "Acima do PL consome capital")
ind("cresc_real", "Crescimento real (descontado o IPCA)", "(1 + nominal) ÷ (1 + IPCA) - 1", lambda a: f"=(1+C{{cresc_carteira}})/(1+{da('ipca', a)})-1", F_PCT1, "Crescimento de verdade")
ind("cresc_comercial", "Crescimento da carteira comercial", "", lambda a: f"={da('s_cred_comercial', a)}/{da('s_cred_comercial', a - 1)}-1", F_PCT1)
ind("cresc_rural", "Crescimento da carteira rural", "", lambda a: f"={da('s_cred_rural', a)}/{da('s_cred_rural', a - 1)}-1", F_PCT1)
ind("cresc_ativo", "Crescimento do ativo", "", lambda a: f"={da('s_ativo', a)}/{da('s_ativo', a - 1)}-1", F_PCT1)
fam("3. CAPTAÇÃO (DEPÓSITOS À VISTA, POUPANÇA, RDC E LCA)")
ind("captacao", "Captação com associados (R$ mil)", "Depósitos + LCA", lambda a: f"={da('s_captacao', a)}", F_MIL, "", "ah")
ind("cresc_captacao", "Crescimento da captação", "", lambda a: f"={da('s_captacao', a)}/{da('s_captacao', a - 1)}-1", F_PCT1, "Se cresce menos que o crédito, a liquidez cai")
ind("mix_dv", "Mix: depósito à vista", "DV ÷ captação", lambda a: f"={bp('dep_vista', a)}/{da('s_captacao', a)}", F_PCT1, "Funding mais barato")
ind("mix_poup", "Mix: poupança", "", lambda a: f"={bp('poupanca', a)}/{da('s_captacao', a)}", F_PCT1)
ind("mix_rdc", "Mix: RDC", "", lambda a: f"={bp('rdc', a)}/{da('s_captacao', a)}", F_PCT1, "Funding mais caro e mais estável")
ind("mix_lca", "Mix: LCA", "", lambda a: f"={bp('lca', a)}/{da('s_captacao', a)}", F_PCT1)
ind("custo_captacao", "Custo médio da captação (inclui DV a custo zero)", "-(poupança + RDC + LCA) da DSP ÷ captação média",
    lambda a: f"=-({ds('desp_poupanca', a)}+{ds('desp_rdc', a)}+{ds('desp_lca', a)})/{media('s_captacao', a)}", F_PCT2)
ind("custo_rdc_cdi", "Custo do RDC em % do CDI", "-(despesa RDC) ÷ RDC médio ÷ CDI", lambda a: f"=-{ds('desp_rdc', a)}/{media('s_rdc', a)}/{da('cdi', a)}", F_PCT1, "Acima de 100% sem crédito bem precificado destrói margem")
ind("credito_captacao", "Crédito ÷ captação", "Carteira ÷ (depósitos + LCA)", lambda a: f"={bp('carteira', a)}/{da('s_captacao', a)}", F_PCT1, "Quanto da captação virou crédito")
fam("4. MARGEM FINANCEIRA")
ind("mf", "Margem financeira (R$ mil)", "Crédito + liquidez - captação - repasses", lambda a: f"={ge('mf', a)}", F_MIL, "", "ah")
ind("nim", "Margem financeira ÷ ativo médio (NIM)", "Margem ÷ média do ativo", lambda a: f"={ge('mf', a)}/{media('s_ativo', a)}", F_PCT2, "Compare sempre com o CDI do período")
ind("taxa_credito", "Taxa média da carteira", "Receita de crédito ÷ carteira média", lambda a: f"={ds('rend_credito', a)}/{media('s_carteira', a)}", F_PCT2)
ind("taxa_liquidez", "Taxa média da liquidez", "Receita da liquidez ÷ liquidez rentável média", lambda a: f"={ge('g_liquidez', a)}/{media('s_liq_rentavel', a)}", F_PCT2, "Perto de 100% do CDI")
fam("5. SPREAD FINANCEIRO")
ind("spread", "Spread de crédito (p.p.)", "Taxa da carteira - custo da captação", lambda a: f"=(C{{taxa_credito}}-C{{custo_captacao}})*100", F_PP, "", "pp_abs")
ind("spread_cdi", "Taxa da carteira - CDI (p.p.)", "Prêmio da carteira sobre o CDI", lambda a: f"=(C{{taxa_credito}}-{da('cdi', a)})*100", F_PP, "Carteira pré perde prêmio quando o CDI sobe", "pp_abs")
fam("6. ROE E RESULTADO RECORRENTE")
ind("resultado", "Resultado antes dos juros ao capital (R$ mil)", "DSP", lambda a: f"={ds('resultado', a)}", F_MIL, "", "ah")
ind("pl_medio", "PL médio (R$ mil)", "Média do PL inicial e final", lambda a: f"={media('s_pl', a)}", F_MIL, "", "ah")
ind("roe", "ROE", "Resultado ÷ PL médio", lambda a: f"=C{{resultado}}/C{{pl_medio}}", F_PCT2, "Compare com o custo de capital (~15%) e com o CDI")
ind("roa", "ROA", "Resultado ÷ ativo médio", lambda a: f"=C{{resultado}}/{media('s_ativo', a)}", F_PCT2)
ind("dp_margem", "DuPont: margem (resultado ÷ receita líquida)", "", lambda a: f"=C{{resultado}}/{ge('receita_liquida', a)}", F_PCT1)
ind("dp_giro", "DuPont: giro (receita líquida ÷ ativo médio)", "", lambda a: f"={ge('receita_liquida', a)}/{media('s_ativo', a)}", F_PCT2)
ind("dp_alav", "DuPont: alavancagem (ativo médio ÷ PL médio)", "", lambda a: f"={media('s_ativo', a)}/C{{pl_medio}}", F_X, "", "x")
ind("nr", "Itens não recorrentes (R$ mil)", "Notas", lambda a: f"={da('nr_total', a)}", F_MIL, "", "ah")
ind("resultado_recorrente", "Resultado recorrente (R$ mil)", "Resultado - não recorrentes", lambda a: f"={ge('resultado_recorrente', a)}", F_MIL, "", "ah")
ind("roe_recorrente", "ROE recorrente", "Resultado recorrente ÷ PL médio", lambda a: f"=C{{resultado_recorrente}}/C{{pl_medio}}", F_PCT2, "O que se repete no ano que vem")
fam("7. EFICIÊNCIA OPERACIONAL (COST-TO-INCOME)")
ind("ie", "Índice de eficiência (IE)", "-Despesas administrativas ÷ receita líquida", lambda a: f"=-{ge('desp_adm', a)}/{ge('receita_liquida', a)}", F_PCT1, "Menor é melhor")
ind("ie_recorrente", "IE recorrente", "Sem PDV na despesa e sem reversão de contingência na receita",
    lambda a: f"=-({ge('desp_adm', a)}-{da('nr_pdv', a)})/({ge('receita_liquida', a)}-{da('nr_reversao_conting', a)})", F_PCT1)
ind("iear", "IE ajustado ao risco (IEAR)", "-(Desp. adm + PDD líquida) ÷ receita líquida", lambda a: f"=-({ge('desp_adm', a)}+{ge('pdd_liquida', a)})/{ge('receita_liquida', a)}", F_PCT1, "Acima de 100% = prejuízo")
fam("8. INADIMPLÊNCIA")
ind("npl90", "Inadimplência acima de 90 dias (NPL 90)", "Vencidos > 90 dias ÷ carteira", lambda a: f"={da('at_d90', a)}/{bp('carteira', a)}", F_PCT2, "Atraso que já virou problema")
ind("inad15_90", "Inadimplência de 15 a 90 dias", "(15-30 + 31-60 + 61-90) ÷ carteira", lambda a: f"=({da('at_d15_30', a)}+{da('at_d31_60', a)}+{da('at_d61_90', a)})/{bp('carteira', a)}", F_PCT2, "Sinal antecedente do NPL 90")
ind("e3_pct", "Estágio 3 ÷ carteira", "Ativos problemáticos", lambda a: f"={da('est_e3', a)}/{bp('carteira', a)}", F_PCT2, "Inclui reestruturados, por isso maior que o NPL 90")
ind("reneg_pct", "Renegociadas ÷ carteira", "", lambda a: f"={da('renegociadas', a)}/{bp('carteira', a)}", F_PCT2, "Renegociação alta pode esconder atraso")
fam("9. PDD / PROVISÃO PARA PERDAS")
ind("prov_carteira", "Provisão ÷ carteira", "-Provisão ÷ carteira", lambda a: f"=-{bp('provisao', a)}/{bp('carteira', a)}", F_PCT2)
ind("cobertura", "Cobertura do NPL 90", "-Provisão ÷ vencidos > 90 dias", lambda a: f"=-{bp('provisao', a)}/{da('at_d90', a)}", F_PCT1, "Abaixo de 100%: provisão não cobre nem o atraso")
ind("cob_e3", "Cobertura do estágio 3", "Provisão do E3 ÷ saldo do E3", lambda a: f"={da('prov_e3', a)}/{da('est_e3', a)}", F_PCT1)
ind("custo_risco", "Custo do risco (PDD ÷ carteira média)", "-PDD ÷ carteira média", lambda a: f"=-{ds('pdd', a)}/{media('s_carteira', a)}", F_PCT2, "Compare com o prêmio da carteira sobre o CDI")
ind("custo_risco_liq", "Custo do risco líquido de recuperações", "-(PDD + recuperação) ÷ carteira média", lambda a: f"=-{ge('pdd_liquida', a)}/{media('s_carteira', a)}", F_PCT2)
ind("baixas", "Baixas para prejuízo (R$ mil)", "Provisão inicial + PDD - provisão final (em módulo)",
    lambda a: f"=-({da('s_provisao', a - 1)}+{ds('pdd', a)}-{da('s_provisao', a)})", F_MIL, "", "ah")
fam("10. RWA (ATIVOS PONDERADOS PELO RISCO)")
ind("rwa", "RWA total (R$ mil)", "Crédito + operacional + mercado", lambda a: f"={rw('rwa', a)}", F_MIL, "", "ah")
ind("rwa_credito", "RWA de crédito (R$ mil)", "Soma de exposição × FPR", lambda a: f"={rw('rwa_credito', a)}", F_MIL, "", "ah")
ind("rwa_operacional", "RWA operacional (R$ mil)", "15% × receita média ÷ 8%", lambda a: f"={rw('rwa_operacional', a)}", F_MIL, "", "ah")
ind("densidade", "Densidade (RWA ÷ ativo)", "", lambda a: f"={rw('densidade', a)}", F_PCT1, "Quanto risco cada real de ativo carrega")
fam("11. ÍNDICE DE BASILEIA")
ind("pr", "Patrimônio de Referência (R$ mil)", "Capital principal + Nível II", lambda a: f"={rw('pr', a)}", F_MIL, "", "ah")
ind("basileia", "Índice de Basileia", "PR ÷ RWA", lambda a: f"={rw('basileia', a)}", F_PCT2, "Mínimo do S4: 10,5%")
ind("mm_rwa", "MM/RWA (folga em p.p.)", "Basileia - 10,5%", lambda a: f"={rw('mm_rwa', a)}", F_PP, "Política de recursos corporativos: piso de 2 p.p.", "pp_abs")
ind("potencial", "Potencial de crédito até o piso (R$ mil)", "(PR ÷ 12,5% - RWA) ÷ 1,1", lambda a: f"={rw('potencial', a)}", F_MIL, "", "ah")
fam("12. LIQUIDEZ (LCR E INDICADORES INTERNOS)")
ind("liquidez", "Liquidez total (R$ mil)", "Caixa + centralização + aplicações + TVM", lambda a: f"={da('s_liquidez', a)}", F_MIL, "", "ah")
ind("liq_ativo", "Liquidez ÷ ativo", "", lambda a: f"={da('s_liquidez', a)}/{bp('ativo', a)}", F_PCT1)
ind("indice_liquidez", "Índice de liquidez interno", "Liquidez ÷ (NML × depósitos)", lambda a: f"={da('s_liquidez', a)}/({par('nml')}*{bp('depositos', a)})", F_X, "Abaixo de 1 falta; muito acima custa", "x")
ind("liq_imediata", "Liquidez imediata", "(Caixa + centralização) ÷ depósitos à vista", lambda a: f"=({bp('caixa', a)}+{bp('centralizacao', a)})/{bp('dep_vista', a)}", F_X, "", "x")
ind("hqla", "LCR: ativos líquidos (R$ mil)", "Caixa + TVM + centralização", lambda a: f"={bp('caixa', a)}+{bp('tvm', a)}+{bp('centralizacao', a)}", F_MIL, "", "ah")
ind("lcr_saidas", "LCR: saídas em 30 dias de estresse (R$ mil)", "Saldos × fatores + vencimentos",
    lambda a: (f"={bp('dep_vista', a)}*{da('lcr_dv_fator', a)}+{bp('poupanca', a)}*{da('lcr_poup_fator', a)}+{bp('rdc', a)}*{da('lcr_rdc_fator', a)}"
               f"+{da('lcr_lca_30d', a)}+{da('lcr_obrig_30d', a)}+{da('lcr_limites', a)}*{da('lcr_limites_fator', a)}+{da('lcr_outras_saidas', a)}"), F_MIL, "", "ah")
ind("lcr_entradas", "LCR: entradas consideradas (R$ mil)", "Mínimo(entradas; 75% das saídas)",
    lambda a: f"=MIN({da('lcr_parcelas_30d', a)}*{da('lcr_parcelas_fator', a)}+{da('lcr_aplic_30d', a)},{par('lcr_teto_entradas')}*C{{lcr_saidas}})", F_MIL, "", "ah")
ind("lcr", "LCR didático", "Ativos líquidos ÷ (saídas - entradas)", lambda a: f"=C{{hqla}}/(C{{lcr_saidas}}-C{{lcr_entradas}})", F_PCT1, "Regulatório só no S1; aqui é régua interna")
fam("13. FTP (FUNDS TRANSFER PRICING): MARGEM POR PRODUTO (R$ mil)")
ind("ftp_comercial", "Crédito comercial: receita - saldo médio × FTP", "", lambda a: f"={da('rend_comercial', a)}-{media('s_cred_comercial', a)}*{da('ftp_comercial', a)}", F_MIL, "", "ah")
ind("ftp_rural", "Crédito rural: receita - saldo médio × FTP", "", lambda a: f"={da('rend_rural', a)}-{media('s_cred_rural', a)}*{da('ftp_rural', a)}", F_MIL, "", "ah")
ind("ftp_dv", "Depósito à vista: saldo médio × FTP - custo (zero)", "", lambda a: f"={media('s_dep_vista', a)}*{da('ftp_dep_vista', a)}", F_MIL, "", "ah")
ind("ftp_poup", "Poupança: saldo médio × FTP - despesa", "", lambda a: f"={media('s_poupanca', a)}*{da('ftp_poupanca', a)}+{ds('desp_poupanca', a)}", F_MIL, "", "ah")
ind("ftp_rdc", "RDC: saldo médio × FTP - despesa", "", lambda a: f"={media('s_rdc', a)}*{da('ftp_rdc', a)}+{ds('desp_rdc', a)}", F_MIL, "", "ah")
ind("ftp_lca", "LCA: saldo médio × FTP - despesa", "", lambda a: f"={media('s_lca', a)}*{da('ftp_lca', a)}+{ds('desp_lca', a)}", F_MIL, "", "ah")
ind("ftp_tesouraria", "Tesouraria/ALM: o que sobra para fechar a margem", "Margem financeira - soma dos produtos",
    lambda a: f"={ge('mf', a)}-SUM(C{{ftp_comercial}}:C{{ftp_lca}})", F_MIL, "Inclui o ganho de financiar ativos com PL e o descasamento", "ah")
fam("CONFERÊNCIAS")
ind("chk_bp", "BP fecha? (ativo - passivo - PL)", "", lambda a: f"={bp('check', a)}", F_MIL, "Deve ser zero", "none")
ind("chk_dsp", "DSP bate com a visão gerencial?", "", lambda a: f"={ge('check', a)}", F_MIL, "Deve ser zero", "none")

r = 5
for x in IND:
    if x[0] == "sec":
        r += 1
        continue
    REG[("Indicadores", x[1])] = r
    r += 1
r = 5
for x in IND:
    if x[0] == "sec":
        ws.cell(row=r, column=1, value=x[1])
        for cc in range(1, 7):
            estilo(ws.cell(row=r, column=cc), COR_TITULO, True, fundo=COR_SECAO)
        r += 1
        continue
    _, k, rot, como, fn, fmt, leitura, var = x
    estilo(ws.cell(row=r, column=1, value=rot))
    estilo(ws.cell(row=r, column=2, value=como), "404040", tam=9)
    for a, col in COL_ANO["Indicadores"].items():
        v = fn(a)
        for kk in re.findall(r"C\{(\w+)\}", v):
            v = v.replace("C{" + kk + "}", f"{col}{REG[('Indicadores', kk)]}")
        estilo(ws.cell(row=r, column="ABCD".index(col) + 1, value=v), cor_de(v), fmt=fmt)
    if var == "pp":
        estilo(ws.cell(row=r, column=5, value=f"=(D{r}-C{r})*100"), PRETO, fmt=F_PP)
    elif var == "pp_abs":
        estilo(ws.cell(row=r, column=5, value=f"=D{r}-C{r}"), PRETO, fmt=F_PP)
    elif var == "x":
        estilo(ws.cell(row=r, column=5, value=f"=D{r}-C{r}"), PRETO, fmt=F_X)
    elif var == "ah":
        estilo(ws.cell(row=r, column=5, value=f'=IFERROR(D{r}/C{r}-1,"n.a.")'), PRETO, fmt=F_PCT1)
    estilo(ws.cell(row=r, column=6, value=leitura), "404040", tam=9)
    r += 1
# RAROC (bloco à parte, só 2025)
r_raroc = REG.get(("Indicadores", "chk_bp")) + 3
ws.cell(row=r_raroc, column=1, value="14. RAROC POR CARTEIRA EM 2025 (R$ mil)")
for cc in range(1, 7):
    estilo(ws.cell(row=r_raroc, column=cc), COR_TITULO, True, fundo=COR_SECAO)
cabecalho(ws, r_raroc + 1, ["Componente", "Como se calcula", "Comercial", "Rural", "", ""])
RAR = [
    ("rr_saldo", "Saldo médio", "Média da carteira", f"={media('s_cred_comercial', 2025)}", f"={media('s_cred_rural', 2025)}", F_MIL),
    ("rr_prov", "Provisão média", "", f"={media('prov_comercial', 2025).replace('$B$', '$C$').replace('$C$', '$C$')}", None, F_MIL),
    ("rr_margem", "Margem após FTP", "Linha de FTP acima", f"=D{REG[('Indicadores', 'ftp_comercial')]}", f"=D{REG[('Indicadores', 'ftp_rural')]}", F_MIL),
    ("rr_custo", "(-) Custo operacional alocado", "Saldo médio × % de custo", "={s}*" + par("raroc_custo_comercial"), "={s}*" + par("raroc_custo_rural"), F_MIL),
    ("rr_pe", "(-) Perda esperada", "Saldo médio × % de PE", "={s}*" + par("raroc_pe_comercial"), "={s}*" + par("raroc_pe_rural"), F_MIL),
    ("rr_rwa", "RWA da carteira", "(Saldo médio - provisão média) × FPR", "=({s}-{p})*" + par("fpr_comercial"), "=({s}-{p})*" + par("fpr_rural"), F_MIL),
    ("rr_capital", "Capital alocado", "RWA × 12,5%", "={rwa}*" + par("raroc_capital"), "={rwa}*" + par("raroc_capital"), F_MIL),
    ("rr_cred_cap", "(+) Crédito do capital (capital × FTP)", "O capital não paga juros: a carteira recebe o FTP sobre ele",
     "={cap}*" + da("ftp_comercial", 2025), "={cap}*" + da("ftp_rural", 2025), F_MIL),
    ("rr_resultado", "(=) Resultado ajustado ao risco", "Margem - custo - PE + crédito do capital", "={m}-{c}-{pe}+{cc}", "={m}-{c}-{pe}+{cc}", F_MIL),
    ("rr_raroc", "RAROC", "Resultado ajustado ÷ capital alocado", "={res}/{cap}", "={res}/{cap}", F_PCT1),
    ("rr_hurdle", "RAROC - hurdle (p.p.)", "Positivo cria valor; negativo destrói", "=({rar}-" + par("raroc_hurdle") + ")*100", "=({rar}-" + par("raroc_hurdle") + ")*100", F_PP),
]
rbase = r_raroc + 2
for n, x in enumerate(RAR):
    REG[("Indicadores", x[0])] = rbase + n
for n, (k, rot, como, fc, fr, fmt) in enumerate(RAR):
    rr = rbase + n
    estilo(ws.cell(row=rr, column=1, value=rot), PRETO, k in ("rr_resultado", "rr_raroc"))
    estilo(ws.cell(row=rr, column=2, value=como), "404040", tam=9)
    for col, f in (("C", fc), ("D", fr)):
        if k == "rr_prov":
            f = f"={media('prov_comercial', 2025)}" if col == "C" else f"={media('prov_rural', 2025)}"
        mp = {"s": f"{col}{REG[('Indicadores', 'rr_saldo')]}", "p": f"{col}{REG[('Indicadores', 'rr_prov')]}",
              "rwa": f"{col}{REG[('Indicadores', 'rr_rwa')]}", "cap": f"{col}{REG[('Indicadores', 'rr_capital')]}",
              "m": f"{col}{REG[('Indicadores', 'rr_margem')]}", "c": f"{col}{REG[('Indicadores', 'rr_custo')]}",
              "pe": f"{col}{REG[('Indicadores', 'rr_pe')]}", "cc": f"{col}{REG[('Indicadores', 'rr_cred_cap')]}",
              "res": f"{col}{REG[('Indicadores', 'rr_resultado')]}", "rar": f"{col}{REG[('Indicadores', 'rr_raroc')]}"}
        for kk, vv in mp.items():
            f = f.replace("{" + kk + "}", vv)
        estilo(ws.cell(row=rr, column="ABCD".index(col) + 1, value=f), cor_de(f), k in ("rr_resultado", "rr_raroc"), fmt)
larguras(ws, [52, 46, 15, 15, 12, 48])
ws.freeze_panes = "C5"

# ================================================================== Exercícios
EXMAP = {e["id"]: e for e in EXERCICIOS}
RESP = {}  # (ex, letra) -> endereço


def folha_exercicio(eid, passos, extra_texto=None):
    e = EXMAP[eid]
    ws = wb.create_sheet("Ex" + eid[1:])
    titulo(ws, f"Exercício {eid[1:]}: {e['titulo']}", f"Tema: {e['tema']} | valores em R$ mil, salvo indicação", 6)
    ws["A3"] = e["enunciado"]
    estilo(ws["A3"], quebra=True)
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 60
    if e.get("dica"):
        ws["A4"] = "Dica: " + e["dica"]
        estilo(ws["A4"], "595959", italico=True, quebra=True)
        ws.merge_cells("A4:F4")
        ws.row_dimensions[4].height = 32
    cabecalho(ws, 6, ["Item", "Passo / pergunta", "Resultado", "Unidade", "Fórmula (Excel PT-BR)", "Como foi feito"])
    loc = {}
    r = 7
    for p in passos:
        loc[p[1]] = r
        r += 1
    r = 7
    for p in passos:
        letra, chave, desc, formula, unidade, expl = p
        f = formula
        if isinstance(f, str):
            for kk in re.findall(r"\{(\w+)\}", f):
                f = f.replace("{" + kk + "}", f"C{loc[kk]}")
        estilo(ws.cell(row=r, column=1, value=letra), PRETO, bool(letra), alinhar="center")
        estilo(ws.cell(row=r, column=2, value=desc), PRETO, bool(letra), quebra=True)
        fmt = FMT_UNIDADE.get(unidade, F_MIL)
        if unidade == "% fator":
            fmt = F_PCT2
        c = ws.cell(row=r, column=3, value=f)
        estilo(c, cor_de(f), bool(letra), fmt, COR_RESPOSTA if letra else None)
        estilo(ws.cell(row=r, column=4, value=unidade.replace(" fator", "")), "595959", tam=9, alinhar="center")
        cf = ws.cell(row=r, column=5, value=pt(f))
        cf.data_type = "s"
        estilo(cf, "404040", tam=9, quebra=True)
        estilo(ws.cell(row=r, column=6, value=expl), PRETO, tam=9, quebra=True)
        ws.row_dimensions[r].height = 42
        if letra:
            RESP[(eid, letra)] = f"Ex{eid[1:]}!$C${r}"
        r += 1
    if extra_texto:
        r += 1
        for t in extra_texto:
            c = ws.cell(row=r, column=1, value=t[0])
            estilo(c, COR_TITULO, True)
            c = ws.cell(row=r, column=2, value=t[1])
            estilo(c, quebra=True)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            ws.row_dimensions[r].height = max(30, 15 * (len(t[1]) // 120 + 1))
            r += 1
    larguras(ws, [6, 52, 16, 9, 52, 70])
    return ws


# ---- Ex01 (classificação)
ws = wb.create_sheet("Ex01")
e = EXMAP["P01"]
titulo(ws, "Exercício 01: " + e["titulo"], "Tema: linhas do BP", 4)
ws["A3"] = e["enunciado"]
estilo(ws["A3"], quebra=True)
ws.merge_cells("A3:D3")
ws.row_dimensions[3].height = 40
cabecalho(ws, 5, ["Nº", "Evento", "Linha do balanço (resposta)", "Por quê"])
for n, (q, resp, expl) in enumerate(CLASSIFICACAO, start=1):
    r = 5 + n
    estilo(ws.cell(row=r, column=1, value=n), alinhar="center")
    estilo(ws.cell(row=r, column=2, value=q), quebra=True)
    estilo(ws.cell(row=r, column=3, value=resp), PRETO, True, fundo=COR_RESPOSTA, quebra=True)
    estilo(ws.cell(row=r, column=4, value=expl), tam=9, quebra=True)
    ws.row_dimensions[r].height = 56
larguras(ws, [5, 60, 34, 90])

A24, A25 = 2024, 2025
fsx = "s_"

folha_exercicio("P02", [
    ("", "tot24", "Total do passivo + PL em 2024", f"={bp('passivo_pl', A24)}", "R$ mil", "Base da AV do lado direito do balanço (igual ao ativo total)."),
    ("", "tot25", "Total do passivo + PL em 2025", f"={bp('passivo_pl', A25)}", "R$ mil", "Mesmo raciocínio para 2025."),
    ("a", "a", "RDC ÷ passivo + PL, 2024", f"={bp('rdc', A24)}/{{tot24}}", "%", "Saldo de RDC de 2024 dividido pelo total."),
    ("b", "b", "RDC ÷ passivo + PL, 2025", f"={bp('rdc', A25)}/{{tot25}}", "%", "Sobe: a cooperativa depende mais do funding que paga % do CDI."),
    ("c", "c", "Depósito à vista ÷ passivo + PL, 2024", f"={bp('dep_vista', A24)}/{{tot24}}", "%", "Funding sem custo."),
    ("d", "d", "Depósito à vista ÷ passivo + PL, 2025", f"={bp('dep_vista', A25)}/{{tot25}}", "%", "Cai: o associado tira dinheiro parado da conta quando o CDI está alto. A captação fica mais cara."),
    ("e", "e", "PL ÷ passivo + PL, 2025", f"={bp('pl', A25)}/{{tot25}}", "%", "Quanto do ativo é financiado com capital próprio (11,3%)."),
    ("f", "f", "(Depósitos + LCA) ÷ passivo + PL, 2025", f"=({bp('depositos', A25)}+{bp('lca', A25)})/{{tot25}}", "%", "Três quartos do balanço vêm dos associados: é o modelo cooperativo."),
], [("Leitura", "O mix piorou de custo: o RDC ganhou espaço e o depósito à vista perdeu. Com CDI alto, isso é esperado, mas pressiona a margem financeira. O PL manteve a participação, ou seja, cresceu no ritmo do balanço.")])

folha_exercicio("P03", [
    ("a", "a", "AH da receita de crédito", f"={ds('rend_credito', A25)}/{ds('rend_credito', A24)}-1", "%", "2025 ÷ 2024 - 1."),
    ("b", "b", "AH da despesa de captação", f"={ds('desp_captacao', A25)}/{ds('desp_captacao', A24)}-1", "%", "Negativo ÷ negativo dá positivo: a fórmula funciona sem ABS."),
    ("c", "c", "AH da margem financeira", f"={ge('mf', A25)}/{ge('mf', A24)}-1", "%", "Margem cresce bem menos que a carteira (14%)."),
    ("d", "d", "AH da despesa de PDD", f"={ds('pdd', A25)}/{ds('pdd', A24)}-1", "%", "PDD cresce quase 3 vezes mais que a carteira: piora de risco."),
    ("e", "e", "AH do resultado", f"={ds('resultado', A25)}/{ds('resultado', A24)}-1", "%", "Resultado cai mesmo com a carteira crescendo."),
], [("Por quê?", "A captação (RDC, LCA, poupança) é paga em % do CDI e reprecifica rápido: com o CDI médio subindo de 10,9% para 14,3%, a despesa explode. A carteira tem parte prefixada (comercial e rural) que só reprecifica no vencimento. A margem cresce pouco e a PDD maior come o resto.")])

folha_exercicio("P04", [
    ("a", "a", "Margem financeira de 2024 (R$ mil)", f"={ds('rend_credito', A24)}+{ds('rend_centralizacao', A24)}+{ds('rend_aplic', A24)}+{ds('rend_tvm', A24)}+{ds('desp_captacao', A24)}+{ds('desp_repasses', A24)}", "R$ mil",
     "Receita de crédito + receita da liquidez + despesa de captação (negativa) + despesa de repasses (negativa)."),
    ("", "am24", "Ativo médio de 2024", f"={media('s_ativo', A24)}", "R$ mil", "Média do ativo de 31/12/2023 (notas) e de 31/12/2024."),
    ("b", "b", "NIM de 2024", "={a}/{am24}", "%", "Margem ÷ ativo médio."),
    ("", "am25", "Ativo médio de 2025", f"={media('s_ativo', A25)}", "R$ mil", ""),
    ("c", "c", "NIM de 2025", f"={ge('mf', A25)}/{{am25}}", "%", "Mesma conta para 2025."),
    ("d", "d", "Variação do NIM (p.p.)", "=({c}-{b})*100", "p.p.", "Queda de meio ponto: a margem cresceu menos que o ativo."),
], [("Leitura", "O NIM caiu porque a captação ficou mais cara mais rápido que a carteira e porque a liquidez (que rende só o CDI) cresceu. Em reais a margem subiu: olhe sempre o percentual.")])

folha_exercicio("P05", [
    ("", "cm", "Carteira média de 2024", f"={media('s_carteira', A24)}", "R$ mil", "Média de 2023 e 2024."),
    ("a", "a", "Taxa média da carteira", f"={ds('rend_credito', A24)}/{{cm}}", "%", "Receita de crédito ÷ carteira média."),
    ("", "capm", "Captação média de 2024 (depósitos + LCA)", f"={media('s_captacao', A24)}", "R$ mil", "Inclui o depósito à vista, que custa zero."),
    ("", "desp", "Despesa de poupança + RDC + LCA", f"=-({ds('desp_poupanca', A24)}+{ds('desp_rdc', A24)}+{ds('desp_lca', A24)})", "R$ mil", "Sem FGCoop (não é juro) e sem LFS (é capital)."),
    ("b", "b", "Custo médio da captação", "={desp}/{capm}", "%", "Despesa ÷ captação média."),
    ("c", "c", "Spread de crédito (p.p.)", "=({a}-{b})*100", "p.p.", "Diferença entre a taxa que recebe e a que paga."),
    ("d", "d", "Custo médio dos passivos onerosos", f"={{desp}}/({{capm}}-{media('s_dep_vista', A24)})", "%", "Tira o depósito à vista do denominador: mostra o custo de quem cobra juros."),
], [("Leitura", "Em 2024 o spread era de 10,3 p.p.; em 2025 cai para 8,6 p.p. O depósito à vista derruba o custo médio (8,5% contra 10,6% só dos onerosos): por isso perder DV dói.")])

folha_exercicio("P06", [
    ("", "res", "Resultado de 2024", f"={ds('resultado', A24)}", "R$ mil", "Antes dos juros ao capital: é o que o capital do associado gerou."),
    ("", "plm", "PL médio de 2024", f"={media('s_pl', A24)}", "R$ mil", "Média de 2023 (notas) e 2024."),
    ("a", "a", "ROE contábil", "={res}/{plm}", "%", "Resultado ÷ PL médio."),
    ("", "nr", "Itens não recorrentes de 2024", f"={da('nr_total', A24)}", "R$ mil", "PDV (-2.500) e venda de imóvel (+300): total -2.200."),
    ("b", "b", "Resultado recorrente", "={res}-{nr}", "R$ mil", "Tirar um item negativo significa somá-lo de volta."),
    ("c", "c", "ROE recorrente", "={b}/{plm}", "%", "Resultado recorrente ÷ PL médio."),
    ("d", "d", "Margem (resultado ÷ receita líquida)", f"={{res}}/{ge('receita_liquida', A24)}", "%", "Quanto de cada real de receita líquida vira resultado."),
    ("e", "e", "Giro (receita líquida ÷ ativo médio)", f"={ge('receita_liquida', A24)}/{media('s_ativo', A24)}", "%", "Quanto de receita líquida cada real de ativo gera."),
    ("f", "f", "Alavancagem (ativo médio ÷ PL médio)", f"={media('s_ativo', A24)}/{{plm}}", "x", "Quantos reais de ativo cada real de PL sustenta."),
    ("", "prova", "Prova: margem × giro × alavancagem", "={d}*{e}*{f}", "%", "Tem de dar o ROE do item a."),
], [("Leitura", "Em 2024 o ROE recorrente (20,7%) era maior que o contábil: o PDV puxou o resultado para baixo. Em 2025 acontece o contrário (16,4% contábil e 13,6% recorrente). Comparar só o contábil esconde metade da piora.")])

folha_exercicio("P07", [
    ("", "rl24", "Receita líquida de 2024", f"={ge('receita_liquida', A24)}", "R$ mil", "Aba Gerencial."),
    ("a", "a", "IE de 2024", f"=-{ge('desp_adm', A24)}/{{rl24}}", "%", "Despesas administrativas (pessoal + outras) ÷ receita líquida."),
    ("b", "b", "IEAR de 2024", f"=-({ge('desp_adm', A24)}+{ge('pdd_liquida', A24)})/{{rl24}}", "%", "Soma a PDD líquida de recuperações."),
    ("", "adm25", "Despesa adm recorrente de 2025", f"=-({ge('desp_adm', A25)}-{da('nr_pdv', A25)})", "R$ mil", "Em 2025 não houve PDV: igual à despesa total."),
    ("", "rl25", "Receita líquida recorrente de 2025", f"={ge('receita_liquida', A25)}-{da('nr_reversao_conting', A25)}", "R$ mil", "Tira a reversão de contingência de 3.000."),
    ("c", "c", "IE recorrente de 2025", "={adm25}/{rl25}", "%", "Pior que o contábil (59,1%)."),
    ("", "alvo", "IE recorrente de 2024 (meta)", f"=-({ge('desp_adm', A24)}-{da('nr_pdv', A24)})/({ge('receita_liquida', A24)}-{da('nr_reversao_conting', A24)})", "%", "Sem o PDV de 2.500."),
    ("d", "d", "Corte necessário (R$ mil)", "={adm25}-{alvo}*{rl25}", "R$ mil", "Despesa atual - despesa que daria o IE meta com a mesma receita."),
], [("Leitura", "O IE contábil quase não mexeu (58,5% para 59,1%), mas o recorrente piorou 2,3 p.p. E o IEAR subiu 5,4 p.p.: o maior problema de 2025 é o risco de crédito, não o custo.")])

folha_exercicio("P08", [
    ("a", "a", "NPL 90 de 2024", f"={da('at_d90', A24)}/{bp('carteira', A24)}", "%", "Vencidos há mais de 90 dias ÷ carteira bruta."),
    ("b", "b", "Inadimplência 15 a 90 dias de 2024", f"=({da('at_d15_30', A24)}+{da('at_d31_60', A24)}+{da('at_d61_90', A24)})/{bp('carteira', A24)}", "%", "Atraso curto que pode virar NPL nos próximos meses."),
    ("c", "c", "Cobertura do NPL 90 em 2024", f"=-{bp('provisao', A24)}/{da('at_d90', A24)}", "%", "Provisão (em módulo) ÷ vencidos > 90 dias."),
    ("d", "d", "Custo do risco de 2024", f"=-{ds('pdd', A24)}/{media('s_carteira', A24)}", "%", "Despesa de PDD ÷ carteira média."),
    ("e", "e", "Baixas para prejuízo de 2024", f"=-({da('s_provisao', 2023)}+{ds('pdd', A24)}-{da('s_provisao', A24)})", "R$ mil", "Provisão inicial (68.000) + PDD (46.000) - provisão final (82.000) = 32.000 baixados."),
    ("f", "f", "Estágio 3 ÷ carteira em 2025", f"={da('est_e3', A25)}/{bp('carteira', A25)}", "%", "Maior que o NPL 90 (2,9%) porque inclui reestruturados e outros sinais de problema."),
], [("Por que a cobertura caiu?", "Em 2025 o NPL 90 quase dobrou (31.000 para 55.000) e a provisão subiu menos (82.000 para 108.000): cobertura de 265% para 196%. Ainda é confortável, mas a tendência pede atenção, junto com as renegociações (de 41.000 para 62.000).")])

folha_exercicio("P09", [
    ("a", "a", "RWA de crédito de 2024", f"={rw('rwa_credito', A24)}", "R$ mil", "Aba RWA: soma de cada exposição × FPR."),
    ("", "rbm", "Receita bruta média 2022 a 2024", f"=AVERAGE({par('receita_bruta_2022')},{par('receita_bruta_2023')},{da('receita_bruta', A24)})", "R$ mil", "Receita bruta = margem financeira + serviços e tarifas."),
    ("b", "b", "RWA operacional de 2024", f"={par('fator_op')}*{{rbm}}/{par('fator_f')}", "R$ mil", "15% × média ÷ 8%."),
    ("", "rwa", "RWA total de 2024", f"={{a}}+{{b}}+{da('rwa_mercado', A24)}", "R$ mil", "Crédito + operacional + mercado."),
    ("", "pr", "PR de 2024", f"={rw('pr', A24)}", "R$ mil", "PL - intangível - outros ajustes + LFS (zero em 2024)."),
    ("c", "c", "Índice de Basileia de 2024", "={pr}/{rwa}", "%", "PR ÷ RWA."),
    ("d", "d", "MM/RWA de 2024", f"=({{c}}-{par('exig_total')})*100", "p.p.", "Folga sobre a exigência de 10,5%."),
    ("e", "e", "Basileia de 2025 sem a LFS", f"={rw('cp', A25)}/{rw('rwa', A25)}", "%", "Capital principal ÷ RWA: sem a LFS o índice teria caído."),
    ("f", "f", "Capital consumido por R$ 200 mi de crédito comercial", f"=200000*{par('fpr_comercial')}*{par('exig_total')}", "R$ mil", "200.000 × 85% × 10,5%."),
    ("", "rwan", "Novo RWA de 2025", f"={rw('rwa', A25)}+200000*{par('fpr_comercial')}-200000*{par('fpr_centralizacao')}", "R$ mil", "Entra crédito com FPR 85% e sai centralização com FPR 20%."),
    ("g", "g", "Novo índice de Basileia de 2025", f"={rw('pr', A25)}/{{rwan}}", "%", "Mesmo PR, RWA maior."),
], [("Leitura", "A cooperativa está acima da régua (MM/RWA perto de 3,9 p.p. em 2025), mas a folga veio da LFS. Crescer R$ 200 mi de crédito comercial tira cerca de 0,7 p.p. do índice. O PL precisa crescer junto com o RWA, e isso depende de resultado retido.")])

folha_exercicio("P10", [
    ("a", "a", "Ativos líquidos de 2024", f"={bp('caixa', A24)}+{bp('tvm', A24)}+{bp('centralizacao', A24)}", "R$ mil", "Caixa + títulos públicos + centralização (resgatável)."),
    ("b", "b", "Saídas em 30 dias de estresse", f"={ref('Indicadores', 'lcr_saidas', 2024)}", "R$ mil",
     "DV × 15% + poupança × 10% + RDC × 12% + LCA e obrigações que vencem + 5% dos limites + outras saídas."),
    ("", "ent", "Entradas consideradas", f"={ref('Indicadores', 'lcr_entradas', 2024)}", "R$ mil", "50% das parcelas a receber + aplicações que vencem, limitado a 75% das saídas."),
    ("c", "c", "LCR de 2024", "={a}/({b}-{ent})", "%", "Ativos líquidos ÷ saídas líquidas."),
    ("d", "d", "Índice de liquidez interno de 2024", f"={da('s_liquidez', A24)}/({par('nml')}*{bp('depositos', A24)})", "x", "Liquidez total ÷ (30% dos depósitos)."),
    ("e", "e", "Crédito ÷ captação de 2024", f"={bp('carteira', A24)}/{da('s_captacao', A24)}", "%", "Carteira ÷ (depósitos + LCA)."),
    ("f", "f", "Liquidez imediata de 2024", f"=({bp('caixa', A24)}+{bp('centralizacao', A24)})/{bp('dep_vista', A24)}", "x", "Quantas vezes o caixa disponível cobre o depósito à vista."),
], [("Leitura", "LCR acima de 400% e índice interno de 1,6: liquidez não é o problema. O ponto é o custo: dinheiro parado rende o CDI e a captação que o financia custa um pouco mais.")])

folha_exercicio("P11", [
    ("", "pm", "Poupança média de 2025", f"={media('s_poupanca', A25)}", "R$ mil", ""),
    ("a", "a", "Margem FTP da poupança", f"={{pm}}*{da('ftp_poupanca', A25)}+{ds('desp_poupanca', A25)}", "R$ mil", "Recebe FTP de 12,0% e paga cerca de 8,3%: margem positiva."),
    ("", "rm", "RDC médio de 2025", f"={media('s_rdc', A25)}", "R$ mil", ""),
    ("b", "b", "Margem FTP do RDC", f"={{rm}}*{da('ftp_rdc', A25)}+{ds('desp_rdc', A25)}", "R$ mil", "Recebe 100% do CDI e paga cerca de 101%: margem negativa."),
    ("c", "c", "Custo do RDC em % do CDI", f"=-{ds('desp_rdc', A25)}/{{rm}}/{da('cdi', A25)}", "%", "Despesa ÷ saldo médio ÷ CDI."),
    ("d", "d", "Custo a mais com RDC a 105% do CDI", f"={{rm}}*{da('cdi', A25)}*1.05+{ds('desp_rdc', A25)}", "R$ mil", "Custo com 105% - custo atual."),
    ("e", "e", "Impacto no ROE (p.p.)", f"=-{{d}}/{media('s_pl', A25)}*100", "p.p.", "Custo a mais ÷ PL médio."),
    ("f", "f", "Resultado da tesouraria (ALM)", f"={ref('Indicadores', 'ftp_tesouraria', 2025)}", "R$ mil", "Margem financeira - soma das margens dos produtos."),
], [("Resposta", "Só vale pagar 105% do CDI se o dinheiro novo for para crédito com taxa que pague esse custo, a provisão e o capital. Se ficar parado na centralização (100% do CDI), cada real captado perde 5% do CDI por ano. O FTP mostra isso produto a produto.")])

folha_exercicio("P12", [
    ("", "v", "Valor da operação", 500, "R$ mil", "Dado do enunciado."),
    ("a", "a", "Perda esperada = valor × PD × LGD", "={v}*3%*45%", "R$ mil", "500 × 3% × 45% = 6,75."),
    ("b", "b", "Capital alocado = valor × FPR × 12,5%", "={v}*100%*12.5%", "R$ mil", "500 × 100% × 12,5% = 62,5."),
    ("", "rec", "Receita de juros (21%)", "={v}*21%", "R$ mil", ""),
    ("", "ftp", "(-) Custo de transferência (FTP 15%)", "=-{v}*15%", "R$ mil", "O que a tesouraria cobra pelo dinheiro."),
    ("", "co", "(-) Custo operacional (2,5%)", "=-{v}*2.5%", "R$ mil", ""),
    ("", "cc", "(+) Crédito do capital (capital × FTP)", "={b}*15%", "R$ mil", "A parte financiada com capital não paga juros."),
    ("c", "c", "Resultado ajustado ao risco", "={rec}+{ftp}+{co}-{a}+{cc}", "R$ mil", "Receita - FTP - custo - PE + crédito do capital."),
    ("d", "d", "RAROC", "={c}/{b}", "%", "Resultado ajustado ÷ capital alocado."),
    ("e", "e", "Taxa mínima para RAROC de 15%", "=(15%*{b}-{ftp}-{co}+{a}-{cc})/{v}", "%", "Isola a taxa: taxa × valor = 15% × capital + FTP + custo + PE - crédito do capital."),
], [("Leitura", "A 21% a operação cria valor (RAROC de 32%). A taxa mínima é 18,85%: abaixo dela, a cooperativa remunera o capital do associado abaixo do custo de capital.")])

folha_exercicio("P13", [
    ("a", "a", "Depósito à vista ÷ captação 2025", f"={bp('dep_vista', A25)}/{da('s_captacao', A25)}", "%", "DV ÷ (depósitos + LCA)."),
    ("b", "b", "RDC ÷ captação 2025", f"={bp('rdc', A25)}/{da('s_captacao', A25)}", "%", ""),
    ("c", "c", "Custo implícito da poupança", f"=-{ds('desp_poupanca', A25)}/{media('s_poupanca', A25)}", "%", "Despesa ÷ saldo médio."),
    ("", "clca", "Custo implícito da LCA", f"=-{ds('desp_lca', A25)}/{media('s_lca', A25)}", "%", ""),
    ("d", "d", "Custo da LCA em % do CDI", f"={{clca}}/{da('cdi', A25)}", "%", "LCA é isenta de IR para PF: paga menos que o CDI."),
    ("e", "e", "Custo anual da migração de R$ 50 mi de DV para RDC", f"=50000*{da('cdi', A25)}*1.01", "R$ mil", "50.000 × 14,3% × 101%: o DV não custava nada."),
], [("Leitura", "Perder depósito à vista é a forma mais silenciosa de perder margem: R$ 50 mi migrando custam R$ 7,2 mi por ano, cerca de 13% do resultado.")])

folha_exercicio("P14", [
    ("a", "a", "Capital social em 31/12/2025", f"=DMPL!$B${REG[('DMPL', 'fim')]}", "R$ mil", "130.000 + 14.000 integralizações - 9.000 devoluções + 12.270 sobras capitalizadas + 12.600 juros ao capital incorporados."),
    ("b", "b", "PL em 31/12/2025", f"=DMPL!$F${REG[('DMPL', 'fim')]}", "R$ mil", "Soma de todas as mutações. Confere com o BP."),
    ("", "capm", "Capital médio de 2025", f"={media('s_capital_social', A25)}", "R$ mil", ""),
    ("c", "c", "Juros ao capital em % do CDI", f"=-{ds('juros_capital', A25)}/({{capm}}*{da('cdi', A25)})", "%", "Juros pagos ÷ (capital médio × CDI)."),
    ("d", "d", "Capital social por associado (R$)", f"={bp('capital_social', A25)}*1000/{da('associados', A25)}", "R$", "R$ mil × 1.000 ÷ número de associados."),
    ("e", "e", "Capital social ÷ PL", f"={bp('capital_social', A25)}/{bp('pl', A25)}", "%", "Parte do PL que pode sair com resgate de quotas."),
    ("f", "f", "Crescimento do PL - crescimento do RWA (p.p.)", f"=({da('s_pl', A25)}/{da('s_pl', A24)}-{rw('rwa', A25)}/{rw('rwa', A24)})*100", "p.p.", "Negativo: o risco cresceu um pouco mais rápido que o capital."),
], [("Leitura", "Quase metade do PL é quota, que pode sair. As reservas são o capital estável. Para sustentar o crescimento, a cooperativa depende de reter sobras; distribuir em conta corrente ajuda o associado hoje, mas não fortalece o capital.")])

# ---- Ex15 (parecer)
ws = wb.create_sheet("Ex15")
titulo(ws, "Exercício 15: Estudo de caso: o parecer da Coop Vale do Ipê", "Resposta modelo. Escreva o seu antes de ler.", 2)
ws["A3"] = EXMAP["P15"]["enunciado"]
estilo(ws["A3"], quebra=True)
ws.merge_cells("A3:B3")
ws.row_dimensions[3].height = 50


from parecer import PARECER

r = 5
for t, txt in PARECER:
    estilo(ws.cell(row=r, column=1, value=t), COR_TITULO, True)
    estilo(ws.cell(row=r, column=2, value=txt), quebra=True)
    ws.row_dimensions[r].height = 62
    r += 1
larguras(ws, [26, 130])

# ================================================================== Gabarito
ws = wb.create_sheet("Gabarito")
titulo(ws, "Gabarito: todas as respostas numéricas", "Cada resposta é um link para a célula da aba do exercício", 5)
cabecalho(ws, 4, ["Exercício", "Item", "Pergunta", "Resposta", "Unidade"])
r = 5
for e in EXERCICIOS:
    for it in e["itens"]:
        if it["unidade"] == "opcao":
            continue
        estilo(ws.cell(row=r, column=1, value=e["id"].replace("P", "Ex")), alinhar="center")
        estilo(ws.cell(row=r, column=2, value=it["letra"]), alinhar="center")
        estilo(ws.cell(row=r, column=3, value=it["pergunta"]), quebra=True)
        v = "=" + RESP[(e["id"], it["letra"])]
        estilo(ws.cell(row=r, column=4, value=v), VERDE, True, FMT_UNIDADE[it["unidade"]], COR_RESPOSTA)
        estilo(ws.cell(row=r, column=5, value=it["unidade"]), "595959", alinhar="center")
        REG[("Gabarito", e["id"] + it["letra"])] = r
        r += 1
larguras(ws, [10, 6, 80, 16, 9])

# ordem das abas
ordem = ["LEIA-ME", "BP", "DSP", "Gerencial", "Dados", "DMPL", "RWA", "Indicadores"] + [f"Ex{n:02d}" for n in range(1, 16)] + ["Gabarito"]
wb._sheets = [wb[n] for n in ordem]
for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(SAIDA)
print("ok", SAIDA)
