"""Confere a planilha recalculada contra o modelo em Python.

Rode depois do recalc (LibreOffice): python3 verificar.py
"""
from openpyxl import load_workbook

import modelo as M
from exercicios import EXERCICIOS

wb = load_workbook("Exercicios_BP_Cooperativa_Resolvidos.xlsx", data_only=True)
erros = []

# 1) respostas do gabarito
g = wb["Gabarito"]
linhas = {}
for row in g.iter_rows(min_row=5, values_only=True):
    if row[0]:
        linhas[(row[0].replace("Ex", "P"), row[1])] = row[3]
n = 0
for e in EXERCICIOS:
    for it in e["itens"]:
        if it["unidade"] == "opcao":
            continue
        v = linhas.get((e["id"], it["letra"]))
        n += 1
        if v is None or abs(v - it["valor"]) > max(1e-6, abs(it["valor"]) * 1e-6):
            erros.append(f"{e['id']}{it['letra']}: planilha={v} python={it['valor']}")
print(f"respostas conferidas: {n}")

# 2) conferências que devem dar zero
for aba, col_ini in (("BP", 1), ("DSP", 1), ("Gerencial", 1), ("DMPL", 1)):
    ws = wb[aba]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] and "Conferência" in str(row[0]):
            for v in row[1:6]:
                if isinstance(v, (int, float)) and abs(v) > 1e-6:
                    erros.append(f"{aba}: conferência diferente de zero ({row[0]}: {v})")

# 3) indicadores principais
ws = wb["Indicadores"]
ind = {}
for row in ws.iter_rows(min_row=5, values_only=True):
    if row[0]:
        ind[row[0]] = row
mapa = {
    "ROE": "roe", "ROE recorrente": "roe_recorrente", "Margem financeira ÷ ativo médio (NIM)": "nim",
    "Índice de eficiência (IE)": "ie", "IE ajustado ao risco (IEAR)": "iear", "Índice de Basileia": "basileia",
    "Inadimplência acima de 90 dias (NPL 90)": "npl90", "LCR didático": "lcr", "RWA total (R$ mil)": "rwa",
    "Custo médio da captação (inclui DV a custo zero)": "custo_captacao", "Índice de liquidez interno": "indice_liquidez",
    "Cobertura do NPL 90": "cobertura_npl", "Crescimento nominal da carteira": "cresc_carteira",
}
for rot, k in mapa.items():
    for a, j in ((2024, 2), (2025, 3)):
        v = ind[rot][j]
        p = M.I[a][k]
        if abs(v - p) > max(1e-6, abs(p) * 1e-6):
            erros.append(f"Indicador {rot} {a}: planilha={v} python={p}")
# RAROC
for row in ws.iter_rows(min_row=5, values_only=True):
    if row[0] == "RAROC":
        for j, seg in ((2, "comercial"), (3, "rural")):
            p = M.I[2025]["raroc"][seg]["raroc"]
            if abs(row[j] - p) > 1e-6:
                erros.append(f"RAROC {seg}: planilha={row[j]} python={p}")
# FTP tesouraria
for a, j in ((2024, 2), (2025, 3)):
    v = ind["Tesouraria/ALM: o que sobra para fechar a margem"][j]
    p = M.I[a]["ftp"]["tesouraria"]
    if abs(v - p) > 1e-3:
        erros.append(f"FTP tesouraria {a}: planilha={v} python={p}")

print("\n".join(erros) if erros else "tudo confere")
raise SystemExit(1 if erros else 0)
